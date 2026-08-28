"""
Сформировать один .xlsx со всеми неудалёнными протоколами (каждый на отдельном листе).

Запуск из каталога backend:

    python scripts/export_all_protocols_xlsx.py
    python scripts/export_all_protocols_xlsx.py -o ../all_protocols.xlsx
    python scripts/export_all_protocols_xlsx.py --laboratory-id 1
    python scripts/export_all_protocols_xlsx.py --laboratory-id 1 --department-id 2
    python scripts/export_all_protocols_xlsx.py --only-with-deleted-methods

Фильтры (опционально):
    --laboratory-id            только протоколы этой лаборатории (все подразделения)
    --department-id            сузить до подразделения (обычно вместе с --laboratory-id)
    --only-with-deleted-methods  только протоколы, где хотя бы один расчёт
                                 связан с удалённым методом исследования
"""

from __future__ import annotations
import argparse
import asyncio
from collections import defaultdict
from copy import copy
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
import re
import sys
from typing import Any, cast
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.ext.asyncio import AsyncSession

BACKEND_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BACKEND_DIR))

from core.config import settings  # noqa: E402
from core.database import AsyncSessionLocal  # noqa: E402
from core.exceptions import BusinessLogicError, ServiceUnavailableError  # noqa: E402
from core.hr_client import hr_post_employees_by_hsnils  # noqa: E402
from core.http_clients import close_all_clients, init_hr_client  # noqa: E402
from core.logger import logger  # noqa: E402
from models.protocol import Protocol  # noqa: E402
from repositories import (  # noqa: E402
    calculation as calculation_repo,
    protocol as protocol_repo,
)
from services.employee import (  # noqa: E402
    build_position_cache_key,
    clear_position_cache,
    prime_position_cache,
)
from services.protocol.generator import generate_protocol_excel  # noqa: E402
from utils.date import ensure_datetime, parse_datetime_string  # noqa: E402
from utils.excel_typing import require_worksheet  # noqa: E402
from utils.hr_employee import (  # noqa: E402
    HR_WORKING_STATUS_ACTIVE,
    format_employee_short_name,
    map_hr_employees_list,
    pick_position_on_date,
)

INVALID_SHEET_NAME_CHARS = re.compile(r"[\\/*?\[\]:]")
MAX_SHEET_NAME_LEN = 31
HR_BATCH_SIZE = 200

_PROTOCOL_SHEET_ERRORS = (
    BusinessLogicError,
    ServiceUnavailableError,
    OSError,
    ValueError,
    TypeError,
    KeyError,
    AttributeError,
)


def resolve_protocol_target_date(protocol: Protocol) -> date | Any:
    """Вернуть дату протокола для запроса должности сотрудника в HR."""
    target_date = protocol.test_protocol_date or protocol.created_at
    if isinstance(target_date, str):
        parsed = ensure_datetime(target_date) or parse_datetime_string(target_date)
        return parsed.date() if parsed is not None else target_date
    if isinstance(target_date, datetime):
        return target_date.date()
    if isinstance(target_date, date):
        return target_date
    parsed = ensure_datetime(target_date)
    return parsed.date() if parsed is not None else target_date


def _merge_employee_records(result: dict[str, dict[str, Any]], employee: dict[str, Any]) -> None:
    """Выбрать актуальную запись сотрудника при дубликатах hsnils."""
    hsnils = employee.get("hsnils")
    if not hsnils:
        return

    if hsnils in result:
        current_status = result[hsnils].get("working_now_status")
        new_status = employee.get("working_now_status")
        if new_status == HR_WORKING_STATUS_ACTIVE and current_status != HR_WORKING_STATUS_ACTIVE:
            result[hsnils] = employee
    else:
        result[hsnils] = employee


async def fetch_employees_by_hashes_post(
    hashes_md5: list[str],
) -> dict[str, dict[str, Any]]:
    """Загрузить один POST-батч сотрудников из HR API с appointments."""
    if not hashes_md5:
        return {}

    if not settings.HR_API_URL:
        raise ValueError("HR_API_URL не настроен")

    status_code, employees_data = await hr_post_employees_by_hsnils(
        hashes_md5,
        include_appointments=True,
    )

    if status_code == 404:
        return {}
    if status_code >= 400:
        raise ServiceUnavailableError(f"HR API недоступен: HTTP {status_code}")

    result: dict[str, dict[str, Any]] = {}
    employees = employees_data if isinstance(employees_data, list) else []
    for employee in map_hr_employees_list([item for item in employees if isinstance(item, dict)]):
        _merge_employee_records(result, employee)

    return result


def resolve_position_and_name(
    employee_data: dict[str, Any],
    target_date: date | datetime | str | Any,
) -> tuple[str, str]:
    """Определить должность и краткое ФИО по уже загруженным данным."""
    full_name = str(employee_data.get("full_name") or "")
    formatted_name = format_employee_short_name(full_name) if full_name else ""
    if not formatted_name:
        return "", ""

    target_datetime = ensure_datetime(target_date)
    if not target_datetime:
        return "", formatted_name

    appointments = employee_data.get("appointments") or []
    if not isinstance(appointments, list) or not appointments:
        return "", formatted_name

    position = pick_position_on_date(
        cast(list[dict[str, Any]], appointments),
        target_datetime,
    )
    return position, formatted_name


def build_position_cache_entries(
    employees_by_hsnils: dict[str, dict[str, Any]],
    position_keys: set[tuple[str, object]],
) -> dict[str, tuple[str, str]]:
    """Собрать записи кэша должностей для generate_protocol_excel."""
    entries: dict[str, tuple[str, str]] = {}
    for hsnils, target_date in position_keys:
        target_datetime = ensure_datetime(target_date)
        if not target_datetime:
            continue
        employee_data = employees_by_hsnils.get(hsnils)
        result = resolve_position_and_name(employee_data, target_date) if employee_data else ("", "")
        entries[build_position_cache_key(hsnils, target_datetime)] = result
    return entries


async def prefetch_hr_for_protocols(db: AsyncSession, protocols: list[Protocol]) -> None:
    """Один раз загрузить сотрудников и должности, нужные для экспорта."""
    clear_position_cache()

    all_hashes: set[str] = set()
    position_keys: set[tuple[str, object]] = set()
    sample_protocol_map: dict[int, list[Protocol]] = defaultdict(list)

    for protocol in protocols:
        target_date = resolve_protocol_target_date(protocol)
        if protocol.issued and not protocol.issued_position:
            all_hashes.add(protocol.issued)
            position_keys.add((protocol.issued, target_date))
        if protocol.approved and not protocol.approved_position:
            all_hashes.add(protocol.approved)
            position_keys.add((protocol.approved, target_date))
        for sample_id in protocol.samples or []:
            sample_protocol_map[sample_id].append(protocol)

    sample_ids = list(sample_protocol_map.keys())
    if sample_ids:
        executor_pairs = await calculation_repo.get_executor_sample_pairs(db, sample_ids)
        for executor, sample_id in executor_pairs:
            for protocol in sample_protocol_map.get(sample_id, []):
                all_hashes.add(executor)
                position_keys.add((executor, resolve_protocol_target_date(protocol)))

    if not all_hashes:
        logger.info("HR: сотрудники в протоколах не найдены")
        return

    logger.info(
        "HR: POST-батч для {} сотрудников, {} пар (hash, дата)",
        len(all_hashes),
        len(position_keys),
    )

    employees_by_hsnils: dict[str, dict[str, Any]] = {}
    hash_list = list(all_hashes)
    for offset in range(0, len(hash_list), HR_BATCH_SIZE):
        chunk = hash_list[offset : offset + HR_BATCH_SIZE]
        employees_by_hsnils.update(await fetch_employees_by_hashes_post(chunk))

    entries = build_position_cache_entries(employees_by_hsnils, position_keys)
    prime_position_cache(entries)
    logger.info(
        "HR: кэш готов ({} сотрудников, {} должностей, только POST-батчи)",
        len(employees_by_hsnils),
        len(entries),
    )


def copy_worksheet(source_ws: Worksheet, target_wb: Workbook, sheet_title: str) -> Worksheet:
    """Скопировать лист протокола в общую книгу с сохранением оформления."""
    target_ws = target_wb.create_sheet(title=sheet_title)

    for row in source_ws.iter_rows():
        for cell in row:
            target_cell = target_ws.cell(row=cell.row, column=cell.column, value=cell.value)
            if cell.has_style:
                target_cell.font = copy(cell.font)
                target_cell.border = copy(cell.border)
                target_cell.fill = copy(cell.fill)
                target_cell.number_format = cell.number_format
                target_cell.protection = copy(cell.protection)
                target_cell.alignment = copy(cell.alignment)

    for col_letter, col_dim in source_ws.column_dimensions.items():
        target_col = target_ws.column_dimensions[col_letter]
        target_col.width = col_dim.width
        target_col.hidden = col_dim.hidden

    for row_idx, row_dim in source_ws.row_dimensions.items():
        target_row = target_ws.row_dimensions[row_idx]
        target_row.height = row_dim.height
        target_row.hidden = row_dim.hidden

    for merged_range in list(source_ws.merged_cells.ranges):
        target_ws.merge_cells(str(merged_range))

    target_ws.page_setup = copy(source_ws.page_setup)
    target_ws.page_margins = copy(source_ws.page_margins)
    target_ws.print_options = copy(source_ws.print_options)
    target_ws.sheet_format = copy(source_ws.sheet_format)
    target_ws.sheet_properties = copy(source_ws.sheet_properties)

    for footer_attr in (
        "oddHeader",
        "oddFooter",
        "evenHeader",
        "evenFooter",
        "firstHeader",
        "firstFooter",
    ):
        if hasattr(source_ws, footer_attr) and hasattr(target_ws, footer_attr):
            setattr(target_ws, footer_attr, copy(getattr(source_ws, footer_attr)))

    return target_ws


def build_sheet_title(protocol: Protocol, used_titles: set[str]) -> str:
    """Собрать уникальное имя листа Excel для протокола."""
    number = (protocol.test_protocol_number or str(protocol.id)).strip()
    number = INVALID_SHEET_NAME_CHARS.sub("_", number)
    number = re.sub(r"\s+", "_", number)
    base = f"{protocol.id}_{number}"[:MAX_SHEET_NAME_LEN]
    if not base:
        base = f"protocol_{protocol.id}"[:MAX_SHEET_NAME_LEN]

    title = base
    suffix = 2
    while title in used_titles:
        tail = f"_{suffix}"
        title = f"{base[: MAX_SHEET_NAME_LEN - len(tail)]}{tail}"
        suffix += 1

    used_titles.add(title)
    return title


async def filter_protocols_with_deleted_methods(
    db: AsyncSession,
    protocols: list[Protocol],
) -> list[Protocol]:
    """Оставить протоколы, у которых есть расчёт с удалённым методом исследования."""
    if not protocols:
        return []

    sample_to_protocol_ids: dict[int, list[int]] = defaultdict(list)
    all_sample_ids: set[int] = set()

    for protocol in protocols:
        for sample_id in protocol.samples or []:
            all_sample_ids.add(sample_id)
            sample_to_protocol_ids[sample_id].append(protocol.id)

    if not all_sample_ids:
        return []

    samples_with_deleted_method = await calculation_repo.get_sample_ids_with_deleted_research_methods(
        db,
        list(all_sample_ids),
    )

    protocol_ids_with_deleted_method: set[int] = set()
    for sample_id in samples_with_deleted_method:
        protocol_ids_with_deleted_method.update(sample_to_protocol_ids.get(sample_id, []))

    return [protocol for protocol in protocols if protocol.id in protocol_ids_with_deleted_method]


def append_protocol_sheet(
    target_wb: Workbook,
    excel_bytes: bytes,
    protocol: Protocol,
    used_titles: set[str],
) -> str:
    """Добавить лист протокола в общую книгу из уже сгенерированных байтов Excel."""
    source_wb = load_workbook(BytesIO(excel_bytes), data_only=False)
    source_ws = require_worksheet(source_wb)
    sheet_title = build_sheet_title(protocol, used_titles)
    copy_worksheet(source_ws, target_wb, sheet_title)
    source_wb.close()
    return sheet_title


def save_export_workbook(target_wb: Workbook, output_path: Path) -> None:
    """Сохранить итоговую книгу экспорта на диск."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    target_wb.save(output_path)
    target_wb.close()


def build_and_save_export_workbook(
    generated: list[tuple[Protocol, bytes]],
    output_path: Path,
) -> tuple[int, list[tuple[int, str]]]:
    """Собрать общую книгу из байтов протоколов и сохранить на диск (sync)."""
    target_wb = Workbook()
    target_wb.remove(require_worksheet(target_wb))

    used_titles: set[str] = set()
    success_count = 0
    skipped: list[tuple[int, str]] = []

    for protocol, excel_bytes in generated:
        protocol_label = protocol.test_protocol_number or protocol.id
        try:
            sheet_title = append_protocol_sheet(target_wb, excel_bytes, protocol, used_titles)
            success_count += 1
            logger.info("OK протокол {} ({}) -> лист «{}»", protocol.id, protocol_label, sheet_title)
        except _PROTOCOL_SHEET_ERRORS as exc:
            message = exc.message if isinstance(exc, BusinessLogicError) else str(exc)
            skipped.append((protocol.id, message))
            logger.warning("Пропуск листа протокола {} ({}): {}", protocol.id, protocol_label, message)

    if success_count == 0:
        target_wb.close()
        return 0, skipped

    save_export_workbook(target_wb, output_path)
    return success_count, skipped


async def generate_protocol_excel_bytes(
    db: AsyncSession,
    protocols: list[Protocol],
) -> tuple[list[tuple[Protocol, bytes]], list[tuple[int, str]]]:
    """Сгенерировать xlsx-байты по протоколам; пропуски — в skipped."""
    generated: list[tuple[Protocol, bytes]] = []
    skipped: list[tuple[int, str]] = []

    for protocol in protocols:
        protocol_label = protocol.test_protocol_number or protocol.id
        try:
            excel_bytes, _filename = await generate_protocol_excel(db, protocol.id)
            generated.append((protocol, excel_bytes))
            logger.info("Сгенерирован протокол {} ({})", protocol.id, protocol_label)
        except _PROTOCOL_SHEET_ERRORS as exc:
            message = exc.message if isinstance(exc, BusinessLogicError) else str(exc)
            skipped.append((protocol.id, message))
            logger.warning("Пропуск протокола {} ({}): {}", protocol.id, protocol_label, message)

    return generated, skipped


async def export_all_protocols(
    laboratory_id: int | None,
    department_id: int | None,
    only_with_deleted_methods: bool = False,
) -> list[tuple[Protocol, bytes]] | None:
    """Подготовить данные и байты протоколов для sync-сборки книги."""
    async with AsyncSessionLocal() as db:
        protocols, _ = await protocol_repo.get_protocols(
            db,
            laboratory_id=laboratory_id,
            department_id=department_id,
            include_deleted=False,
            sort_by="id",
            sort_order="asc",
        )

        if only_with_deleted_methods:
            total_before = len(protocols)
            protocols = await filter_protocols_with_deleted_methods(db, protocols)
            logger.info(
                "Фильтр --only-with-deleted-methods: {} из {} протоколов",
                len(protocols),
                total_before,
            )

        if not protocols:
            if only_with_deleted_methods:
                logger.warning("Не найдено протоколов с удалёнными методами исследования")
            else:
                logger.warning("Не найдено неудалённых протоколов")
            return None

        filter_parts: list[str] = []
        if laboratory_id is not None:
            filter_parts.append(f"лаборатория={laboratory_id}")
        if department_id is not None:
            filter_parts.append(f"подразделение={department_id}")
        if only_with_deleted_methods:
            filter_parts.append("только с удалёнными методами")
        if filter_parts:
            logger.info("Фильтр: {}", ", ".join(filter_parts))
        else:
            logger.info("Фильтр: все лаборатории и подразделения")

        logger.info("Найдено протоколов: {}", len(protocols))

        init_hr_client()
        try:
            await prefetch_hr_for_protocols(db, protocols)
            generated, skipped = await generate_protocol_excel_bytes(db, protocols)
            if skipped:
                logger.warning("Пропущено при генерации: {}", len(skipped))
                for protocol_id, reason in skipped:
                    logger.warning("  id={}: {}", protocol_id, reason)
            return generated if generated else None
        finally:
            clear_position_cache()
            await close_all_clients()


def parse_args() -> argparse.Namespace:
    """Разобрать аргументы командной строки."""
    parser = argparse.ArgumentParser(description="Экспорт всех неудалённых протоколов в один .xlsx (лист на протокол).")
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        default=BACKEND_DIR / "scripts" / "all_protocols.xlsx",
        help="Путь к результирующему файлу .xlsx",
    )
    parser.add_argument(
        "--laboratory-id",
        type=int,
        default=None,
        metavar="ID",
        help=(
            "ID лаборатории: экспорт протоколов только этой лаборатории "
            "(по всем подразделениям, если --department-id не указан)"
        ),
    )
    parser.add_argument(
        "--department-id",
        type=int,
        default=None,
        metavar="ID",
        help="ID подразделения: дополнительно сузить выборку (обычно вместе с --laboratory-id)",
    )
    parser.add_argument(
        "--only-with-deleted-methods",
        action="store_true",
        help=(
            "Экспортировать только протоколы, у которых хотя бы один неудалённый расчёт "
            "связан с удалённым методом исследования"
        ),
    )
    return parser.parse_args()


def main() -> None:
    """Точка входа: async-подготовка данных, sync-сборка xlsx."""
    args = parse_args()
    output_path = args.output.resolve()
    generated = asyncio.run(
        export_all_protocols(
            laboratory_id=args.laboratory_id,
            department_id=args.department_id,
            only_with_deleted_methods=args.only_with_deleted_methods,
        )
    )
    if not generated:
        logger.error("Ни один протокол не сформирован, файл не создан")
        return

    success_count, skipped = build_and_save_export_workbook(generated, output_path)
    if success_count == 0:
        logger.error("Ни один протокол не сформирован, файл не создан")
        if skipped:
            logger.error("Пропущено протоколов: {}", len(skipped))
        return

    logger.info("Сохранено: {}", output_path)
    logger.info("Листов: {} из {}", success_count, len(generated))
    if skipped:
        logger.warning("Пропущено протоколов: {}", len(skipped))
        for protocol_id, reason in skipped:
            logger.warning("  id={}: {}", protocol_id, reason)


if __name__ == "__main__":
    main()
