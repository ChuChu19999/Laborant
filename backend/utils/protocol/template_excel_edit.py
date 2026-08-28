from __future__ import annotations
from typing import Any
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet
from utils.excel_typing import resolve_style_cell, set_cell_value
from utils.protocol.template_markers import copy_cell_style, get_template_content_bounds

HEADER_MARKERS_MISSING = (
    "В файле не найдены метки {{start_header}} и {{end_header}}. Добавьте метки в шаблон для редактирования шапки."
)


def find_header_marker_rows(worksheet: Worksheet) -> tuple[int | None, int | None]:
    """Ищет {{start_header}} / {{end_header}} в колонке A в пределах содержимого листа."""
    start_header_row: int | None = None
    end_header_row: int | None = None
    last_row, _ = get_template_content_bounds(worksheet)
    for row_idx in range(1, last_row + 1):
        value = worksheet.cell(row=row_idx, column=1).value
        if value is None:
            continue
        if value == "{{start_header}}":
            start_header_row = row_idx
        elif value == "{{end_header}}":
            end_header_row = row_idx
            break
    return start_header_row, end_header_row


def _font_weight_label(font: Font) -> str:
    """Определить начертание по флагу bold и имени шрифта."""
    if font.bold is True or getattr(font, "b", None) in (True, 1, "1", "true"):
        return "bold"
    name = (font.name or "").casefold()
    if any(token in name for token in ("bold", "полужир", "semibold", "black")):
        return "bold"
    return "normal"


def _font_style_label(font: Font) -> str:
    """Определить курсив по флагу italic и имени шрифта."""
    if font.italic is True or getattr(font, "i", None) in (True, 1, "1", "true"):
        return "italic"
    name = (font.name or "").casefold()
    if any(token in name for token in ("italic", "курсив", "oblique")):
        return "italic"
    return "normal"


def _font_size_px(font: Font) -> str:
    """Вернуть размер шрифта в CSS-формате px."""
    size = font.size if font.size is not None else 14
    if isinstance(size, float) and size.is_integer():
        size = int(size)
    return f"{size}px"


def _text_align_label(cell) -> str:
    """Вернуть выравнивание текста ячейки."""
    if cell.alignment and cell.alignment.horizontal == "left":
        return "left"
    if cell.alignment and cell.alignment.horizontal == "right":
        return "right"
    return "center"


def _cell_style_dict(cell) -> dict[str, str]:
    """Собрать словарь стиля ячейки для редактора шапки."""
    font = cell.font or Font()
    style: dict[str, str] = {
        "font_weight": _font_weight_label(font),
        "font_style": _font_style_label(font),
        "font_size": _font_size_px(font),
        "text_align": _text_align_label(cell),
    }
    if font.name:
        style["font_family"] = font.name
    return style


def extract_header_cell_styles(worksheet: Worksheet) -> dict[str, dict[str, str]]:
    """
    Собирает стили ячеек колонки A между метками шапки.

    Ключи стиля — поля приложения: font_weight, font_style, font_size, text_align, font_family.
    """
    start_header_row, end_header_row = find_header_marker_rows(worksheet)
    if start_header_row is None or end_header_row is None:
        raise ValueError(HEADER_MARKERS_MISSING)

    styles: dict[str, dict[str, str]] = {}
    for row_idx in range(start_header_row + 1, end_header_row):
        cell = resolve_style_cell(worksheet, row_idx, 1)
        cell_key = f"{row_idx - start_header_row - 1}-0"
        styles[cell_key] = _cell_style_dict(cell)
    return styles


def _parse_font_size(raw: Any) -> int:
    """Разобрать размер шрифта из CSS-подобного значения (px, число); иначе 14."""
    if isinstance(raw, str):
        raw = raw.replace("px", "")
        try:
            return int(float(raw))
        except (ValueError, TypeError):
            return 14
    if isinstance(raw, (int, float)):
        return int(raw)
    return 14


def _style_value(style: dict[str, Any], snake: str, camel: str, default: Any) -> Any:
    """Прочитать поле стиля: snake_case (канон бэкенда) или camelCase (как с фронта)."""
    if snake in style:
        return style[snake]
    if camel in style:
        return style[camel]
    return default


def apply_header_section_edits(
    worksheet: Worksheet,
    data: list[list[Any]],
    styles: dict[str, dict[str, Any]],
) -> None:
    """Перезаписывает область шапки между метками: сдвиг строк, значения и стили."""
    start_header_row, end_header_row = find_header_marker_rows(worksheet)
    if start_header_row is None or end_header_row is None:
        raise ValueError(HEADER_MARKERS_MISSING)

    _, template_last_col = get_template_content_bounds(worksheet)

    if end_header_row - start_header_row - 1 < len(data):
        shift = len(data) - (end_header_row - start_header_row - 1)
        for row_idx in range(worksheet.max_row, end_header_row - 1, -1):
            for col_idx in range(1, template_last_col + 1):
                source_cell = worksheet.cell(row=row_idx, column=col_idx)
                target_cell = worksheet.cell(row=row_idx + shift, column=col_idx)

                if isinstance(source_cell, MergedCell):
                    for merge_range in list(worksheet.merged_cells.ranges):
                        if (
                            merge_range.min_row <= row_idx <= merge_range.max_row
                            and merge_range.min_col <= col_idx <= merge_range.max_col
                        ):
                            source_cell = worksheet.cell(
                                row=merge_range.min_row,
                                column=merge_range.min_col,
                            )
                            worksheet.merge_cells(
                                start_row=merge_range.min_row + shift,
                                start_column=merge_range.min_col,
                                end_row=merge_range.max_row + shift,
                                end_column=merge_range.max_col,
                            )
                            break

                if not isinstance(target_cell, MergedCell) and not isinstance(source_cell, MergedCell):
                    target_cell.value = source_cell.value
                    copy_cell_style(source_cell, target_cell)
        end_header_row += shift

    merged_ranges: list[dict[str, int]] = []
    for merge_range in worksheet.merged_cells.ranges:
        if start_header_row < merge_range.min_row < end_header_row:
            merged_ranges.append(
                {
                    "min_row": merge_range.min_row,
                    "max_row": merge_range.max_row,
                    "min_col": merge_range.min_col,
                    "max_col": merge_range.max_col,
                }
            )

    for row_idx in range(start_header_row + 1, end_header_row):
        ranges_to_remove = [
            merge_range
            for merge_range in worksheet.merged_cells.ranges
            if merge_range.min_row <= row_idx <= merge_range.max_row
        ]
        for merge_range in ranges_to_remove:
            worksheet.unmerge_cells(
                start_row=merge_range.min_row,
                start_column=merge_range.min_col,
                end_row=merge_range.max_row,
                end_column=merge_range.max_col,
            )
        set_cell_value(worksheet, row_idx, 1, None)

    for idx, row_data in enumerate(data):
        value = str(row_data[0]) if row_data and row_data[0] is not None else ""
        target_row = start_header_row + 1 + idx
        cell = set_cell_value(worksheet, target_row, 1, value)
        if cell is None:
            continue

        cell_key = f"{idx}-0"
        style = styles.get(cell_key)
        if not style:
            continue

        font_size = _parse_font_size(_style_value(style, "font_size", "fontSize", "14"))
        font_name = _style_value(style, "font_family", "fontFamily", "Times New Roman")
        cell.font = Font(
            name=font_name,
            bold=_style_value(style, "font_weight", "fontWeight", "normal") == "bold",
            italic=_style_value(style, "font_style", "fontStyle", "normal") == "italic",
            size=font_size,
            color="FF000000",
        )

        text_align = _style_value(style, "text_align", "textAlign", "center")
        horizontal_align = "center"
        if text_align == "left":
            horizontal_align = "left"
        elif text_align == "right":
            horizontal_align = "right"

        cell.alignment = Alignment(
            horizontal=horizontal_align,
            vertical="center",
            wrap_text=True,
        )

    for merge_info in merged_ranges:
        worksheet.merge_cells(
            start_row=merge_info["min_row"],
            start_column=merge_info["min_col"],
            end_row=merge_info["max_row"],
            end_column=merge_info["max_col"],
        )
