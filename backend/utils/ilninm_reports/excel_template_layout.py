from __future__ import annotations
import contextlib
from openpyxl.worksheet.worksheet import Worksheet


def insert_rows_for_data_count(
    ws: Worksheet,
    template_data_row: int,
    data_row_count: int,
) -> None:
    """Сдвигает хвост шаблона отчёта вниз под строки данных."""
    if data_row_count <= 1:
        return
    ws.insert_rows(template_data_row + 1, data_row_count - 1)


def unmerge_cells_in_row_range(
    ws: Worksheet,
    first_row: int,
    last_row: int,
    max_col: int,
) -> None:
    """Снимает объединения ячеек только в диапазоне строк данных отчёта."""
    to_unmerge: list[str] = []
    for merged in list(ws.merged_cells.ranges):
        if merged.max_row < first_row or merged.min_row > last_row:
            continue
        if merged.min_col > max_col or merged.max_col < 1:
            continue
        to_unmerge.append(str(merged))
    for range_str in to_unmerge:
        with contextlib.suppress(KeyError):
            ws.unmerge_cells(range_str)
