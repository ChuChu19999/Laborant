from __future__ import annotations
from typing import Any, cast
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet


def require_worksheet(workbook: Workbook, title: str | None = None) -> Worksheet:
    """Вернуть лист книги; падает, если листа нет."""
    sheet = workbook.active if title is None else workbook[title]
    if not isinstance(sheet, Worksheet):
        raise TypeError(f"Ожидался Worksheet, получено {type(sheet)!r}")
    return sheet


def set_cell_value(worksheet: Worksheet, row: int, column: int, value: Any) -> Cell | None:
    """Записать значение в ячейку, пропуская MergedCell."""
    cell = worksheet.cell(row=row, column=column)
    if isinstance(cell, MergedCell):
        return None
    cell.value = value
    return cell


def as_writable_cell(cell: Cell | MergedCell) -> Cell | None:
    """Вернуть обычную Cell или None для MergedCell."""
    if isinstance(cell, MergedCell):
        return None
    return cast(Cell, cell)


def resolve_style_cell(worksheet: Worksheet, row: int, column: int) -> Cell:
    """Вернуть ячейку-владельца стиля: для MergedCell — верхний левый угол диапазона."""
    cell = worksheet.cell(row=row, column=column)
    if not isinstance(cell, MergedCell):
        return cell

    for merge_range in worksheet.merged_cells.ranges:
        if merge_range.min_row <= row <= merge_range.max_row and merge_range.min_col <= column <= merge_range.max_col:
            return cast(
                Cell,
                worksheet.cell(row=merge_range.min_row, column=merge_range.min_col),
            )

    resolved = worksheet.cell(row=row, column=column)
    if isinstance(resolved, MergedCell):
        raise TypeError(f"Не удалось разрешить стиль ячейки ({row}, {column})")
    return resolved
