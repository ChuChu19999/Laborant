import base64
from copy import copy
from io import BytesIO
from typing import Any, Dict
import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Font
from sqlalchemy.ext.asyncio import AsyncSession
from core.exceptions import NotFoundError, ValidationError
from core.logger import logger
from models.protocol import ProtocolTemplate
from repositories import protocol as protocol_repo
from repositories.base import flush_entity
from services.protocol import get_protocol_template_by_id
from utils.versioning import next_version_string


async def get_template_file(
    template: ProtocolTemplate, section: str | None = None
) -> bytes:
    """Получить файл шаблона в виде байтов."""
    file_data = template.file
    try:
        template_bytes = BytesIO(base64.b64decode(file_data))
    except Exception:
        try:
            with open(file_data, "rb") as f:
                template_bytes = BytesIO(f.read())
        except Exception:
            template_bytes = BytesIO(
                file_data.encode() if isinstance(file_data, str) else file_data
            )

    return template_bytes.getvalue()


async def get_excel_styles(
    db: AsyncSession, template_id: int, section: str
) -> Dict[str, Any]:
    """Получить стили для ячеек в файле."""
    template = await get_protocol_template_by_id(db, template_id)
    if not template:
        raise NotFoundError("Шаблон протокола не найден")

    file_data = template.file
    try:
        template_bytes = BytesIO(base64.b64decode(file_data))
    except Exception:
        try:
            with open(file_data, "rb") as f:
                template_bytes = BytesIO(f.read())
        except Exception:
            template_bytes = BytesIO(
                file_data.encode() if isinstance(file_data, str) else file_data
            )

    workbook = openpyxl.load_workbook(template_bytes)
    worksheet = workbook.active

    styles: Dict[str, Any] = {}

    # Ищем метки в файле
    start_header_row = None
    end_header_row = None
    for row_idx in range(1, worksheet.max_row + 1):
        cell_value = worksheet.cell(row=row_idx, column=1).value
        if cell_value == "{{start_header}}":
            start_header_row = row_idx
        elif cell_value == "{{end_header}}":
            end_header_row = row_idx
            break

    # Проверяем наличие меток
    if start_header_row is None or end_header_row is None:
        error_message = (
            "В файле не найдены метки {{start_header}} и {{end_header}}. "
            "Добавьте метки в шаблон для редактирования шапки."
        )
        logger.error(error_message)
        return {"error": error_message}

    # Получаем стили для ячеек между метками
    for row_idx in range(start_header_row + 1, end_header_row):
        cell = worksheet.cell(row=row_idx, column=1)
        cell_key = f"{row_idx - start_header_row - 1}-0"

        if cell.font:
            font_size = cell.font.size or 14
            styles[cell_key] = {
                "fontWeight": "bold" if cell.font.bold else "normal",
                "fontStyle": "italic" if cell.font.italic else "normal",
                "fontSize": f"{font_size}px",
            }

            if cell.alignment:
                if cell.alignment.horizontal == "left":
                    styles[cell_key]["textAlign"] = "left"
                elif cell.alignment.horizontal == "right":
                    styles[cell_key]["textAlign"] = "right"
                else:
                    styles[cell_key]["textAlign"] = "center"
            else:
                styles[cell_key]["textAlign"] = "center"
        else:
            styles[cell_key] = {
                "fontWeight": "normal",
                "fontStyle": "normal",
                "fontSize": "14px",
                "textAlign": "center",
            }

    return {"styles": styles}


async def save_excel_section(
    db: AsyncSession,
    template_id: int,
    data: list,
    styles: Dict[str, Any],
    section: str,
) -> Dict[str, Any]:
    """Сохранить изменения в секции Excel файла."""
    current_template = await get_protocol_template_by_id(db, template_id)
    if not current_template:
        raise NotFoundError("Шаблон протокола не найден")

    # Деактивируем текущий шаблон и получаем следующую версию
    latest_template = await protocol_repo.get_latest_protocol_template(
        db,
        current_template.name,
        current_template.laboratory_id,
        current_template.department_id,
    )
    next_version = next_version_string(
        latest_template.version if latest_template else None
    )

    # Помечаем текущий шаблон как удаленный
    current_template.soft_delete()
    await flush_entity(db)

    file_data = current_template.file
    try:
        template_bytes = BytesIO(base64.b64decode(file_data))
    except Exception:
        try:
            with open(file_data, "rb") as f:
                template_bytes = BytesIO(f.read())
        except Exception:
            template_bytes = BytesIO(
                file_data.encode() if isinstance(file_data, str) else file_data
            )

    workbook = openpyxl.load_workbook(template_bytes)
    worksheet = workbook.active

    # В зависимости от типа секции применяем логику
    if section == "header":
        # Находим существующие метки в файле
        start_header_row = None
        end_header_row = None
        for row_idx in range(1, worksheet.max_row + 1):
            cell_value = worksheet.cell(row=row_idx, column=1).value
            if cell_value == "{{start_header}}":
                start_header_row = row_idx
            elif cell_value == "{{end_header}}":
                end_header_row = row_idx

        # Проверяем наличие меток
        if start_header_row is None or end_header_row is None:
            error_message = (
                "В файле не найдены метки {{start_header}} и {{end_header}}. "
                "Добавьте метки в шаблон для редактирования шапки."
            )
            logger.error(error_message)
            raise ValidationError(error_message)

        # Если нужно больше места между метками
        if end_header_row - start_header_row - 1 < len(data):
            # Сдвигаем данные после end_header вниз
            shift = len(data) - (end_header_row - start_header_row - 1)
            for row_idx in range(worksheet.max_row, end_header_row - 1, -1):
                for col_idx in range(1, worksheet.max_column + 1):
                    source_cell = worksheet.cell(row=row_idx, column=col_idx)
                    target_cell = worksheet.cell(row=row_idx + shift, column=col_idx)

                    # Если исходная ячейка объединена, находим основную ячейку
                    if isinstance(source_cell, MergedCell):
                        for merge_range in worksheet.merged_cells.ranges:
                            min_row = merge_range.min_row
                            max_row = merge_range.max_row
                            min_col = merge_range.min_col
                            max_col = merge_range.max_col

                            if (
                                min_row <= row_idx <= max_row
                                and min_col <= col_idx <= max_col
                            ):
                                source_cell = worksheet.cell(
                                    row=min_row, column=min_col
                                )
                                # Создаем новый диапазон объединения со сдвигом
                                worksheet.merge_cells(
                                    start_row=min_row + shift,
                                    start_column=min_col,
                                    end_row=max_row + shift,
                                    end_column=max_col,
                                )
                                break

                    # Копируем значение и стили
                    if not isinstance(target_cell, MergedCell):
                        target_cell.value = source_cell.value
                        if source_cell.has_style:
                            target_cell._style = copy(source_cell._style)
            end_header_row += shift

        # Сохраняем информацию об объединенных ячейках
        merged_ranges = []
        for merge_range in worksheet.merged_cells.ranges:
            if start_header_row < merge_range.min_row < end_header_row:
                merged_ranges.append(
                    {
                        "min_row": merge_range.min_row,
                        "max_row": merge_range.max_row,
                        "min_col": merge_range.min_col,
                        "max_col": merge_range.max_col,
                    }
                )

        # Очищаем старые данные между метками и разъединяем ячейки
        for row_idx in range(start_header_row + 1, end_header_row):
            # Находим и удаляем объединения ячеек в этой области
            ranges_to_remove = []
            for merge_range in worksheet.merged_cells.ranges:
                min_row = merge_range.min_row
                max_row = merge_range.max_row

                if min_row <= row_idx <= max_row:
                    ranges_to_remove.append(merge_range)

            for merge_range in ranges_to_remove:
                worksheet.unmerge_cells(
                    start_row=merge_range.min_row,
                    start_column=merge_range.min_col,
                    end_row=merge_range.max_row,
                    end_column=merge_range.max_col,
                )

            worksheet.cell(row=row_idx, column=1).value = None

        # Записываем новые данные
        for idx, row_data in enumerate(data):
            value = str(row_data[0]) if row_data[0] is not None else ""
            target_row = start_header_row + 1 + idx
            target_col = 1
            cell = worksheet.cell(row=target_row, column=target_col)
            cell.value = value

            # Применяем стили, если они есть для данной ячейки
            cell_key = f"{idx}-0"
            if styles and cell_key in styles:
                style = styles[cell_key]

                font_size = style.get("fontSize", "14")
                if isinstance(font_size, str):
                    font_size = font_size.replace("px", "")
                    try:
                        font_size = int(float(font_size))
                    except (ValueError, TypeError):
                        font_size = 14

                font = Font(
                    name="Times New Roman",
                    bold=style.get("fontWeight") == "bold",
                    italic=style.get("fontStyle") == "italic",
                    size=font_size,
                )
                cell.font = font

                horizontal_align = "center"
                if style.get("textAlign") == "left":
                    horizontal_align = "left"
                elif style.get("textAlign") == "right":
                    horizontal_align = "right"

                alignment = Alignment(
                    horizontal=horizontal_align,
                    vertical="center",
                    wrap_text=True,
                )
                cell.alignment = alignment

        # Восстанавливаем объединенные ячейки
        for merge_info in merged_ranges:
            worksheet.merge_cells(
                start_row=merge_info["min_row"],
                start_column=merge_info["min_col"],
                end_row=merge_info["max_row"],
                end_column=merge_info["max_col"],
            )

    else:
        logger.error(f"Неизвестная секция для редактирования: {section}")
        raise ValidationError(f"Неизвестная секция: {section}")

    output = BytesIO()
    workbook.save(output)
    file_content = output.getvalue()
    file_base64 = base64.b64encode(file_content).decode("utf-8")

    # Создаем новый шаблон с измененным содержимым
    new_template = ProtocolTemplate(
        name=current_template.name,
        version=next_version,
        file=file_base64,
        file_name=current_template.file_name,
        laboratory_id=current_template.laboratory_id,
        department_id=current_template.department_id,
        accreditation_header_row=current_template.accreditation_header_row,
    )
    new_template = await protocol_repo.add_protocol_template(db, new_template)
    await flush_entity(db)

    return {
        "message": "Файл успешно обновлен",
        "version": next_version,
        "template_id": new_template.id,
    }
