import re
from datetime import date, datetime
from io import BytesIO
from typing import Any
import pendulum
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.ext.asyncio import AsyncSession
from schemas.sample_export import SampleExportItem
from services.calculation_export_format import format_sample_calculations_column
from services.employees import get_employees_by_hsnils
from services.sample_export import get_samples_export_data
from utils.sample_formatting import format_well_display

EXPORT_HEADERS = [
    "№ пробы",
    "Тип пробы",
    "Объект испытания",
    "Филиал",
    "Место отбора",
    "Дата отбора",
    "Дата получения",
    "Протоколы",
    "Дата создания",
    "Добавил пробу",
    "Условия отбора",
    "Методы и результаты",
    "Количество показателей",
]

COLUMN_WIDTHS = [14, 24, 28, 20, 32, 14, 14, 28, 14, 24, 24, 48, 12]

EXPORT_FONT = Font(name="Times New Roman", size=12, color="FF000000")
THIN_SIDE = Side(style="thin", color="000000")
EXPORT_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
EXPORT_ALIGNMENT = Alignment(vertical="top", wrap_text=True)


def _format_date(value: Any) -> str:
    if not value:
        return "-"
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    try:
        return pendulum.parse(str(value)).format("DD.MM.YYYY")
    except Exception:
        return str(value)


def _format_sampling_location(sample: SampleExportItem) -> str:
    parts: list[str] = []
    if sample.sampling_location_name:
        parts.append(sample.sampling_location_name)
    well_display = format_well_display(sample.well)
    if well_display:
        parts.append(well_display)
    if sample.mode:
        parts.append(sample.mode)
    return " ".join(parts) if parts else "-"


def _format_protocols(protocols: list[dict[str, Any]] | None) -> str:
    if not protocols:
        return "-"

    formatted = [
        str(protocol.get("formatted_protocol_number") or "-")
        for protocol in protocols
        if protocol.get("formatted_protocol_number")
    ]
    return ", ".join(formatted) if formatted else "-"


def _format_selection_conditions(
    conditions: dict[str, Any] | None,
) -> str:
    if not conditions:
        return "-"

    lines: list[str] = []
    for variable, value in conditions.items():
        if value is None or str(value).strip() == "":
            continue
        formatted_value = re.sub(r"(-?\d+)\.(\d+)", r"\1,\2", str(value))
        lines.append(f"{variable} = {formatted_value}")

    return "\n".join(lines) if lines else "-"


def _apply_cell_style(cell) -> None:
    cell.font = EXPORT_FONT
    cell.border = EXPORT_BORDER
    cell.alignment = EXPORT_ALIGNMENT


def _build_sample_row(
    sample: SampleExportItem,
    employees_map: dict[str, dict[str, Any]],
) -> list[Any]:
    added_by_name = "-"
    if sample.added_by:
        employee = employees_map.get(sample.added_by, {})
        added_by_name = employee.get("fullName") or "-"

    calculations = sample.calculations or []

    return [
        sample.registration_number or "-",
        sample.sample_type or "-",
        sample.test_object or "-",
        sample.branch_name or "-",
        _format_sampling_location(sample),
        _format_date(sample.sampling_date),
        _format_date(sample.receiving_date),
        _format_protocols(sample.protocols),
        _format_date(sample.created_at),
        added_by_name,
        _format_selection_conditions(sample.selection_conditions),
        format_sample_calculations_column(calculations),
        sample.indicators_count,
    ]


def build_samples_export_workbook(
    items: list[SampleExportItem],
    employees_map: dict[str, dict[str, Any]],
) -> bytes:
    """Собирает xlsx-файл таблицы поступления проб."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Поступления проб"

    rows = [EXPORT_HEADERS] + [
        _build_sample_row(sample, employees_map) for sample in items
    ]

    for row_index, row_values in enumerate(rows, start=1):
        for col_index, value in enumerate(row_values, start=1):
            cell = worksheet.cell(row=row_index, column=col_index, value=value)
            _apply_cell_style(cell)

    for col_index, width in enumerate(COLUMN_WIDTHS, start=1):
        worksheet.column_dimensions[get_column_letter(col_index)].width = width

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


async def build_samples_export_excel(
    db: AsyncSession,
    laboratory_id: int | None = None,
    department_id: int | None = None,
    search: str | None = None,
    search_sampling_location: str | None = None,
    search_protocols: str | None = None,
    search_added_by: str | None = None,
    sample_type: str | None = None,
    sample_types: list[str] | None = None,
    test_object: str | None = None,
    test_objects: list[str] | None = None,
    sort_by: str | None = None,
    sort_order: str | None = None,
    sampling_date_from: pendulum.DateTime | None = None,
    sampling_date_to: pendulum.DateTime | None = None,
    receiving_date_from: pendulum.DateTime | None = None,
    receiving_date_to: pendulum.DateTime | None = None,
    created_at_from: pendulum.DateTime | None = None,
    created_at_to: pendulum.DateTime | None = None,
) -> tuple[bytes, int]:
    """Возвращает байты xlsx и количество строк данных."""
    items, total = await get_samples_export_data(
        db,
        laboratory_id=laboratory_id,
        department_id=department_id,
        search=search,
        search_sampling_location=search_sampling_location,
        search_protocols=search_protocols,
        search_added_by=search_added_by,
        sample_type=sample_type,
        sample_types=sample_types,
        test_object=test_object,
        test_objects=test_objects,
        sort_by=sort_by,
        sort_order=sort_order,
        sampling_date_from=sampling_date_from,
        sampling_date_to=sampling_date_to,
        receiving_date_from=receiving_date_from,
        receiving_date_to=receiving_date_to,
        created_at_from=created_at_from,
        created_at_to=created_at_to,
    )

    if not items:
        return b"", 0

    hsnils_list = list({sample.added_by for sample in items if sample.added_by})
    employees_map: dict[str, dict[str, Any]] = {}
    if hsnils_list:
        try:
            employees_map = await get_employees_by_hsnils(
                hsnils_list, include_photo=False
            )
        except Exception:
            employees_map = {}

    excel_bytes = build_samples_export_workbook(items, employees_map)
    return excel_bytes, total
