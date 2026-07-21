import base64
import re
from contextvars import ContextVar
from copy import copy
from io import BytesIO
from typing import Any, Dict, List, Optional
from urllib.parse import quote
import openpyxl
import orjson
import pendulum
from fastapi import HTTPException, status
from fastapi.responses import Response
from openpyxl.styles import Border
from openpyxl.worksheet.header_footer import _HeaderFooterPart
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload
from core.exceptions import NotFoundError, ValidationError
from core.logger import logger
from models.calculation import Calculation
from models.equipment import Equipment
from models.protocol import Protocol
from models.research import ResearchMethod
from models.sample import Sample, SelectionConditions
from services.employees import (
    get_employee_position_and_name,
)
from services.test_object import (
    get_protocol_abbreviations_by_names,
    pick_first_protocol_abbreviation,
)
from utils.protocol_generator_utils import (
    adjust_cell_height_if_needed,
    apply_sheet_print_area,
    check_method_name,
    copy_cell_style,
    copy_column_dimensions,
    copy_row_formatting,
    copy_row_with_styles,
    copy_sheet_page_settings,
    find_protocol_end_row,
    format_measurement_error_value,
    format_protocol_calculation_result,
    get_template_content_bounds,
    group_name_matches,
    join_unique_values,
    parse_if_line_condition,
    strip_if_line_marker,
    template_contains_marker,
)

# Аббревиатура для текущего формирования Excel (без протягивания по всем функциям).
_protocol_abbreviation_ctx: ContextVar[str] = ContextVar(
    "protocol_abbreviation", default=""
)

TABLE1_END_MARKERS = {
    "{end_table1}",
}
TABLE_START_MARKERS = {
    "{start_table1}",
    "{start_table2}",
    "{start_table3}",
}
TABLE_END_MARKERS = {
    "{end_table1}",
    "{end_table2}",
    "{end_table3}",
}


async def process_cell_markers(
    protocol: Protocol,
    samples: List[Sample],
    cell_value: str,
    selection_conditions_templates: Optional[List[Dict[str, Any]]] = None,
    sampling_location_name_only: bool = False,
) -> str:
    """Обрабатывает все метки в ячейке."""
    if not cell_value or not isinstance(cell_value, str):
        return cell_value

    if "{" not in cell_value:
        return cell_value

    try:
        if "{accreditation}" in cell_value:
            if not protocol.is_accredited:
                return "HIDE_ROW"
            cell_value = cell_value.replace("{accreditation}", "").strip()
            if "{" not in cell_value:
                return cell_value

        # Не аккредитован — после номера протокола в этой ячейке ничего не выводим.
        if not protocol.is_accredited and "{test_protocol_number}" in cell_value:
            prefix = cell_value.split("{test_protocol_number}", 1)[0]
            cell_value = f"{prefix}{{test_protocol_number}}"

        result = cell_value
        start = 0
        while True:
            start = result.find("{", start)
            if start == -1:
                break

            end = result.find("}", start)
            if end == -1:
                break

            marker = result[start + 1 : end]
            if marker.startswith("if line"):
                start = end + 1
                continue
            if not marker.startswith("sel_cond_") and marker != "bu":
                value = await get_marker_value_title(
                    protocol,
                    samples,
                    marker,
                    sampling_location_name_only=sampling_location_name_only,
                )
                result = result.replace(f"{{{marker}}}", value)
                continue

            start = end + 1

        processed_result = process_selection_conditions_row(
            samples, result, selection_conditions_templates
        )
        if processed_result is None:
            return "HIDE_ROW"
        return processed_result

    except Exception as e:
        logger.error(f"Ошибка при обработке меток в ячейке: {str(e)}")
        return cell_value


async def get_marker_value_title(
    protocol: Protocol,
    samples: List[Sample],
    marker: str,
    *,
    sampling_location_name_only: bool = False,
) -> str:
    """Возвращает значение для метки в заголовке протокола."""
    try:
        if marker == "test_protocol_number":
            return str(protocol.test_protocol_number or "").strip()

        if marker == "date_protocol":
            if not protocol.test_protocol_date:
                return ""
            return pendulum.instance(protocol.test_protocol_date).format("DD.MM.YYYY")

        if marker == "abbreviation":
            return _protocol_abbreviation_ctx.get() or ""

        if marker == "accreditation":
            return ""

        elif marker == "subd":
            branches = [
                sample.branch.name
                for sample in samples
                if sample.branch and sample.branch.name
            ]
            return join_unique_values(branches)

        elif marker == "tel":
            phones = [sample.phone for sample in samples if sample.phone]
            return join_unique_values(phones)

        elif marker == "res_object":
            objects = [sample.test_object for sample in samples if sample.test_object]
            return join_unique_values(objects)

        elif marker == "sampling_location":
            locations = []
            for sample in samples:
                location_parts = []
                if sample.sampling_location and sample.sampling_location.name:
                    location_parts.append(sample.sampling_location.name.strip())

                if not sampling_location_name_only:
                    if sample.well and sample.well.strip():
                        well_value = sample.well.strip()
                        if re.match(r"^\d", well_value):
                            well_value = f"скв. {well_value}"
                        location_parts.append(well_value)

                    if sample.mode and sample.mode.strip():
                        location_parts.append(sample.mode.strip())

                if location_parts:
                    locations.append(" ".join(location_parts))
            return join_unique_values(locations)

        elif marker == "mode":
            modes = [
                sample.mode.strip()
                for sample in samples
                if sample.mode and sample.mode.strip()
            ]
            return join_unique_values(modes)

        elif marker == "sampling_date":
            dates = sorted(
                set(
                    pendulum.instance(sample.sampling_date).format("DD.MM.YYYY")
                    for sample in samples
                    if sample.sampling_date
                )
            )
            return ", ".join(dates) if dates else ""

        elif marker == "receiving_date":
            dates = sorted(
                set(
                    pendulum.instance(sample.receiving_date).format("DD.MM.YYYY")
                    for sample in samples
                    if sample.receiving_date
                )
            )
            return ", ".join(dates) if dates else ""

        elif marker == "laboratory_activity_dates":
            dates = []
            for sample in samples:
                for calc in sample.calculations:
                    if calc.deleted_at is None and calc.laboratory_activity_date:
                        dates.append(calc.laboratory_activity_date)
            if dates:
                min_date = pendulum.instance(min(dates)).format("DD.MM.YYYY")
                max_date = pendulum.instance(max(dates)).format("DD.MM.YYYY")
                return f"{min_date}-{max_date}" if min_date != max_date else min_date
            return ""

        elif marker == "lab_location":
            if (
                protocol.department
                and hasattr(protocol.department, "laboratory_location")
                and protocol.department.laboratory_location
            ):
                return protocol.department.laboratory_location
            elif (
                protocol.laboratory
                and hasattr(protocol.laboratory, "laboratory_location")
                and protocol.laboratory.laboratory_location
            ):
                return protocol.laboratory.laboratory_location
            return ""

        elif marker == "sampling_act_number":
            return protocol.sampling_act_number or ""

        elif marker == "registration_number":
            numbers = [
                sample.registration_number
                for sample in samples
                if sample.registration_number
            ]
            return join_unique_values(numbers)

        elif marker == "workplace_issued":
            if protocol.issued_position:
                return protocol.issued_position
            elif protocol.issued:
                target_date = protocol.test_protocol_date or protocol.created_at
                if isinstance(target_date, str):
                    target_date = pendulum.parse(target_date)
                elif hasattr(target_date, "date"):
                    target_date = target_date.date()
                position, _ = await get_employee_position_and_name(
                    protocol.issued, target_date
                )
                return position
            return ""

        elif marker == "issued":
            if protocol.issued:
                target_date = protocol.test_protocol_date or protocol.created_at
                if isinstance(target_date, str):
                    target_date = pendulum.parse(target_date)
                elif hasattr(target_date, "date"):
                    target_date = target_date.date()
                _, formatted_name = await get_employee_position_and_name(
                    protocol.issued, target_date
                )
                return formatted_name
            return ""

        elif marker == "workplace_approved":
            if protocol.approved_position:
                return protocol.approved_position
            elif protocol.approved:
                target_date = protocol.test_protocol_date or protocol.created_at
                if isinstance(target_date, str):
                    target_date = pendulum.parse(target_date)
                elif hasattr(target_date, "date"):
                    target_date = target_date.date()
                position, _ = await get_employee_position_and_name(
                    protocol.approved, target_date
                )
                return position
            return ""

        elif marker == "approved":
            if protocol.approved:
                target_date = protocol.test_protocol_date or protocol.created_at
                if isinstance(target_date, str):
                    target_date = pendulum.parse(target_date)
                elif hasattr(target_date, "date"):
                    target_date = target_date.date()
                _, formatted_name = await get_employee_position_and_name(
                    protocol.approved, target_date
                )
                return formatted_name
            return ""

        return ""
    except Exception as e:
        logger.error(f"Ошибка при получении значения для метки {marker}: {str(e)}")
        return ""


def process_selection_conditions_row(
    samples: List[Sample],
    cell_value: str,
    selection_conditions_templates: Optional[List[Dict[str, Any]]] = None,
) -> Optional[str]:
    """Обрабатывает метки условий отбора в ячейке."""
    if not cell_value or not isinstance(cell_value, str):
        return cell_value

    if "{sel_cond_" not in cell_value and "{bu}" not in cell_value:
        return cell_value

    # Создаем словарь для поиска единиц измерения по названию переменной
    unit_map = {}
    if selection_conditions_templates:
        for template in selection_conditions_templates:
            if isinstance(template, dict) and "conditions" in template:
                conditions = template.get("conditions") or []
                if isinstance(conditions, list):
                    for condition_template in conditions:
                        if isinstance(condition_template, dict):
                            variable = condition_template.get("variable", "")
                            unit = condition_template.get("unit", "")
                            if variable:
                                unit_map[variable] = unit

    all_conditions = []
    for sample in samples:
        if not sample.selection_conditions:
            continue

        sample_conditions = sample.selection_conditions

        # Поддержка формата, когда условия хранятся в виде
        # {"conditions": [ ... ]}
        if isinstance(sample_conditions, dict) and "conditions" in sample_conditions:
            sample_conditions = sample_conditions.get("conditions") or []

        # Поддержка формата, когда условия хранятся в виде
        # {"Давление": "4.33", "Температура": "-6", ...}
        # где ключ - это название переменной, значение - это значение
        # Единицы измерения берутся из шаблонов SelectionConditions
        if (
            isinstance(sample_conditions, dict)
            and "conditions" not in sample_conditions
        ):
            for variable, value in sample_conditions.items():
                if variable and value and str(value).strip() != "":
                    formatted_value = str(value).replace(".", ",")
                    # Ищем единицу измерения в шаблонах по названию переменной
                    unit = unit_map.get(variable, "")
                    all_conditions.append(
                        {
                            "variable": variable,
                            "value": formatted_value,
                            "unit": unit,
                        }
                    )
            continue

        if isinstance(sample_conditions, list):
            for condition_item in sample_conditions:
                if isinstance(condition_item, dict):
                    variable = condition_item.get("variable", "")
                    value = condition_item.get("value", "")
                    unit = condition_item.get("unit", "")

                    if variable and value and str(value).strip() != "":
                        formatted_value = str(value).replace(".", ",")
                        all_conditions.append(
                            {
                                "variable": variable,
                                "value": formatted_value,
                                "unit": unit,
                            }
                        )

    has_condition_tags = "{sel_cond_" in cell_value

    if not all_conditions:
        if has_condition_tags:
            return None
        if "{bu}" in cell_value:
            return cell_value.replace("{bu}", "б/у")
        return cell_value

    current_index = None
    for i in range(1, 6):
        if re.search(r"\{sel_cond_[^}]*_" + str(i) + r"\}", cell_value):
            current_index = i
            break

    if has_condition_tags and current_index is None:
        return None

    if current_index is not None and current_index > len(all_conditions):
        return None

    if current_index is not None and current_index <= len(all_conditions):
        condition = all_conditions[current_index - 1]

        if not condition["value"] or condition["value"].strip() == "":
            return None

        result = cell_value
        name_aliases = {"name", "var", "variable"}
        value_aliases = {"val", "value"}
        unit_aliases = {"unit", "units", "measure"}

        tags_for_index = re.findall(
            r"\{sel_cond_([^}]+)_" + str(current_index) + r"\}", result
        )
        for tag_key in set(tags_for_index):
            if tag_key in name_aliases:
                result = result.replace(
                    f"{{sel_cond_{tag_key}_{current_index}}}",
                    condition["variable"],
                )
            elif tag_key in value_aliases:
                result = result.replace(
                    f"{{sel_cond_{tag_key}_{current_index}}}",
                    condition["value"],
                )
            elif tag_key in unit_aliases:
                result = result.replace(
                    f"{{sel_cond_{tag_key}_{current_index}}}",
                    condition["unit"],
                )
            else:
                result = result.replace(f"{{sel_cond_{tag_key}_{current_index}}}", "")

        if "{bu}" in result:
            result = result.replace("{bu}", "")

        if result.strip() == "" or result.strip() == " ":
            return None

        return result

    if "{bu}" in cell_value:
        cell_value = cell_value.replace("{bu}", "")

    return cell_value


async def process_header(
    protocol: Protocol,
    samples: List[Sample],
    template_sheet,
    new_sheet,
    merged_cells_map,
    selection_conditions_templates: Optional[List[Dict[str, Any]]] = None,
    sampling_location_name_only: bool = False,
):
    """Обрабатывает шапку протокола."""
    current_row_new = 1
    found_start = False

    for current_row in range(1, template_sheet.max_row + 1):
        row = list(
            template_sheet.iter_rows(
                min_row=current_row, max_row=current_row, values_only=True
            )
        )[0]

        if not found_start:
            if any(cell and str(cell).strip() == "{{start_header}}" for cell in row):
                found_start = True
            continue

        if any(cell and str(cell).strip() == "{{end_header}}" for cell in row):
            return current_row + 1

        if current_row in template_sheet.row_dimensions:
            new_sheet.row_dimensions[current_row_new] = copy(
                template_sheet.row_dimensions[current_row]
            )

        for merged_range in template_sheet.merged_cells.ranges:
            if merged_range.min_row == current_row:
                new_range = openpyxl.worksheet.cell_range.CellRange(
                    min_col=merged_range.min_col,
                    min_row=current_row_new,
                    max_col=merged_range.max_col,
                    max_row=current_row_new
                    + (merged_range.max_row - merged_range.min_row),
                )
                merged_cells_map.add(new_range)

        copy_row_with_styles(template_sheet, new_sheet, current_row, current_row_new)

        skip_row = False
        for col in range(1, template_sheet.max_column + 1):
            cell = new_sheet.cell(row=current_row_new, column=col)
            if cell.value:
                processed_value = await process_cell_markers(
                    protocol,
                    samples,
                    str(cell.value),
                    selection_conditions_templates,
                    sampling_location_name_only,
                )
                if processed_value == "HIDE_ROW":
                    skip_row = True
                    break
                elif processed_value is not None:
                    cell.value = processed_value

        if skip_row:
            new_sheet.row_dimensions[current_row_new].hidden = True

        current_row_new += 1

    return current_row


async def process_header_and_conditions(
    protocol: Protocol,
    samples: List[Sample],
    template_sheet,
    new_sheet,
    start_row,
    merged_cells_map,
    selection_conditions_templates: Optional[List[Dict[str, Any]]] = None,
    sampling_location_name_only: bool = False,
):
    """Обрабатывает заголовок и условия отбора после шапки до начала таблицы."""
    current_row_new = new_sheet.max_row + 1
    template_last_row, _ = get_template_content_bounds(template_sheet)

    for current_row in range(start_row, template_last_row + 1):
        row = list(
            template_sheet.iter_rows(
                min_row=current_row, max_row=current_row, values_only=True
            )
        )[0]

        if any(cell and str(cell).strip() == "{start_table1}" for cell in row):
            return current_row

        skip_row = False
        processed_values = []

        for cell_value in row:
            if cell_value:
                processed_value = await process_cell_markers(
                    protocol,
                    samples,
                    str(cell_value),
                    selection_conditions_templates,
                    sampling_location_name_only,
                )
                if processed_value == "HIDE_ROW":
                    skip_row = True
                    break
                elif processed_value is not None:
                    processed_values.append(processed_value)
                else:
                    processed_values.append(cell_value)
            else:
                processed_values.append(cell_value)

        if skip_row:
            continue

        if current_row in template_sheet.row_dimensions:
            new_sheet.row_dimensions[current_row_new] = copy(
                template_sheet.row_dimensions[current_row]
            )

        for merged_range in template_sheet.merged_cells.ranges:
            if merged_range.min_row == current_row:
                new_range = openpyxl.worksheet.cell_range.CellRange(
                    min_col=merged_range.min_col,
                    min_row=current_row_new,
                    max_col=merged_range.max_col,
                    max_row=current_row_new
                    + (merged_range.max_row - merged_range.min_row),
                )
                merged_cells_map.add(new_range)

        copy_row_with_styles(template_sheet, new_sheet, current_row, current_row_new)

        for col, processed_value in enumerate(processed_values, start=1):
            if processed_value is not None:
                cell = new_sheet.cell(row=current_row_new, column=col)
                cell.value = processed_value

        current_row_new += 1

    return current_row


async def process_footer_test_protocol_number(
    protocol: Protocol, samples: List[Sample], text
) -> _HeaderFooterPart:
    """Подставляет номер протокола и аббревиатуру в тексте колонтитула."""
    orig_text = text
    if hasattr(text, "text"):
        text = text.text
    if not text or not isinstance(text, str):
        return orig_text

    # Не аккредитован — после номера протокола ничего не выводим.
    if not protocol.is_accredited and "{test_protocol_number}" in text:
        prefix = text.split("{test_protocol_number}", 1)[0]
        text = f"{prefix}{{test_protocol_number}}"

    number = await get_marker_value_title(protocol, samples, "test_protocol_number")
    abbreviation = await get_marker_value_title(protocol, samples, "abbreviation")
    date_protocol = await get_marker_value_title(protocol, samples, "date_protocol")
    result = text.replace("{test_protocol_number}", number)
    result = result.replace("{abbreviation}", abbreviation)
    result = result.replace("{date_protocol}", date_protocol)

    prefix = '&"Times New Roman,Обычный"'
    size_mark = "&11"
    if result.startswith(prefix):
        after = result[len(prefix) :]
        if after.startswith("&") and len(after) > 1 and after[1].isdigit():
            idx = 2
            while idx < len(after) and after[idx].isdigit():
                idx += 1
            after = after[idx:]
        if not after.startswith(size_mark):
            after = size_mark + after
        result = prefix + after
    else:
        result = f"{prefix}{size_mark}{result}"

    return _HeaderFooterPart(text=result)


async def process_footer(
    protocol: Protocol,
    samples: List[Sample],
    template_sheet,
    current_sheet,
    footer_start,
    merged_cells_map,
    current_row,
    selection_conditions_templates: Optional[List[Dict[str, Any]]] = None,
    sampling_location_name_only: bool = False,
):
    """Обрабатывает оставшиеся строки после последней таблицы (подвал протокола)."""
    template_last_row, template_last_col = get_template_content_bounds(template_sheet)

    for row_num in range(footer_start + 1, template_last_row + 1):
        row = list(
            template_sheet.iter_rows(min_row=row_num, max_row=row_num, values_only=True)
        )[0]
        if any(cell and str(cell).strip() in TABLE_END_MARKERS for cell in row):
            continue

        if row_num in template_sheet.row_dimensions:
            current_sheet.row_dimensions[current_row] = copy(
                template_sheet.row_dimensions[row_num]
            )

        for merged_range in template_sheet.merged_cells.ranges:
            if merged_range.min_row == row_num:
                new_range = openpyxl.worksheet.cell_range.CellRange(
                    min_col=merged_range.min_col,
                    min_row=current_row,
                    max_col=merged_range.max_col,
                    max_row=current_row + (merged_range.max_row - merged_range.min_row),
                )
                merged_cells_map.add(new_range)

        for col in range(1, template_last_col + 1):
            source_cell = template_sheet.cell(row=row_num, column=col)
            target_cell = current_sheet.cell(row=current_row, column=col)

            target_cell.value = source_cell.value
            copy_cell_style(source_cell, target_cell)

        for col in range(1, template_last_col + 1):
            cell = current_sheet.cell(row=current_row, column=col)
            if cell.value:
                processed_value = await process_cell_markers(
                    protocol,
                    samples,
                    str(cell.value),
                    selection_conditions_templates,
                    sampling_location_name_only,
                )
                if processed_value is not None:
                    cell.value = processed_value

        if find_protocol_end_row(current_sheet) == current_row:
            current_row += 1
            break

        current_row += 1

    return current_sheet


async def process_between_tables(
    protocol: Protocol,
    samples: List[Sample],
    template_sheet,
    current_sheet,
    table_end,
    merged_cells_map,
    current_row,
    selection_conditions_templates: Optional[List[Dict[str, Any]]] = None,
    sampling_location_name_only: bool = False,
):
    """Обрабатывает данные между таблицами."""
    template_last_row, _ = get_template_content_bounds(template_sheet)

    next_table_start = None
    for row_num in range(table_end + 1, template_last_row + 1):
        row = list(
            template_sheet.iter_rows(min_row=row_num, max_row=row_num, values_only=True)
        )[0]
        if any(cell and str(cell).strip() in TABLE_START_MARKERS for cell in row):
            next_table_start = row_num
            break

    if not next_table_start:
        return current_sheet, current_row

    executors_cache = set()
    target_date = protocol.test_protocol_date or protocol.created_at
    if hasattr(target_date, "date"):
        target_date = target_date.date()
    elif isinstance(target_date, str):
        target_date = pendulum.parse(target_date).date()

    unique_executors = set()
    for sample in samples:
        for calc in sample.calculations:
            if calc.deleted_at is None and calc.executor:
                unique_executors.add(calc.executor)

    for executor_hsnils in unique_executors:
        position, formatted_name = await get_employee_position_and_name(
            executor_hsnils, target_date
        )
        if formatted_name:
            logger.info(
                f"Из HR API получили: ФИО={formatted_name}, должность={position or ''}"
            )
            position_lower = position.lower() if position else ""
            executor_info = (
                f"{position_lower} {formatted_name}".strip()
                if position_lower
                else formatted_name
            )
            executors_cache.add(executor_info)

    executors = sorted(executors_cache) if executors_cache else []

    row_with_executor = None
    executor_column = None

    for row_num in range(table_end + 1, next_table_start):
        row = list(
            template_sheet.iter_rows(min_row=row_num, max_row=row_num, values_only=True)
        )[0]

        if any(cell and str(cell).strip() in TABLE_END_MARKERS for cell in row):
            continue

        for col_idx, cell_value in enumerate(row, 1):
            if cell_value and "{executor}" in str(cell_value):
                row_with_executor = row_num
                executor_column = col_idx
                break

        if row_with_executor:
            break

    for row_num in range(table_end + 1, next_table_start):
        row = list(
            template_sheet.iter_rows(min_row=row_num, max_row=row_num, values_only=True)
        )[0]

        if any(cell and str(cell).strip() in TABLE_END_MARKERS for cell in row):
            continue

        copy_row_formatting(
            template_sheet, current_sheet, row_num, current_row, merged_cells_map
        )

        if row_num == row_with_executor and executors and executor_column:
            for col in range(1, template_sheet.max_column + 1):
                cell = current_sheet.cell(row=current_row, column=col)
                if cell.value:
                    if col == executor_column and "{executor}" in str(cell.value):
                        cell.value = str(cell.value).replace("{executor}", executors[0])
                    else:
                        processed_value = await process_cell_markers(
                            protocol,
                            samples,
                            str(cell.value),
                            selection_conditions_templates,
                            sampling_location_name_only,
                        )
                        if processed_value is not None:
                            cell.value = processed_value

            current_row += 1

            for executor in executors[1:]:
                copy_row_formatting(
                    template_sheet,
                    current_sheet,
                    row_num,
                    current_row,
                    merged_cells_map,
                )

                for col in range(1, template_sheet.max_column + 1):
                    cell = current_sheet.cell(row=current_row, column=col)
                    if cell.value:
                        if col == executor_column and "{executor}" in str(cell.value):
                            cell.value = str(cell.value).replace("{executor}", executor)
                        else:
                            cell.value = ""

                current_row += 1
        else:
            for col in range(1, template_sheet.max_column + 1):
                cell = current_sheet.cell(row=current_row, column=col)
                if cell.value:
                    processed_value = await process_cell_markers(
                        protocol,
                        samples,
                        str(cell.value),
                        selection_conditions_templates,
                        sampling_location_name_only,
                    )
                    if processed_value is not None:
                        cell.value = processed_value

            current_row += 1

    return current_sheet, current_row


def _method_group_names(method: ResearchMethod) -> list[str]:
    """Активные имена групп методики."""
    names: list[str] = []
    for group in method.groups or []:
        if group.deleted_at is not None:
            continue
        name = (group.name or "").strip()
        if name:
            names.append(name)
    return names


def _primary_group_name(method: ResearchMethod | None) -> str:
    """Первое активное имя группы методики или пустая строка."""
    if not method:
        return ""
    names = _method_group_names(method)
    return names[0] if names else ""


def _find_calculation_for_if_line(
    conditions: dict[str, str],
    calculations: List[Calculation],
) -> Calculation | None:
    """Находит расчёт, подходящий под условие {if line}."""
    required_name = (conditions.get("name_method") or "").strip()
    required_group = (conditions.get("group_name") or "").strip()

    for calc in calculations:
        method = calc.research_method
        if not method:
            continue
        if required_name and (method.name or "").strip() != required_name:
            continue
        if required_group:
            group_names = _method_group_names(method)
            if not any(
                group_name_matches(required_group, name) for name in group_names
            ):
                continue
        return calc
    return None


def _row_has_fractional_markers(template_sheet, row_num: int) -> bool:
    """Проверяет, есть ли в строке шаблона fractional-метки."""
    for col in range(1, template_sheet.max_column + 1):
        value = template_sheet.cell(row=row_num, column=col).value
        if value and isinstance(value, str) and "{fractional_" in value:
            return True
    return False


def _fill_method_row_placeholders(
    current_sheet,
    row_num: int,
    *,
    method_id: str | None,
    calc: Calculation | None,
    clear_all: bool = False,
) -> None:
    """Подставляет метки таблицы 1."""
    for col in range(1, current_sheet.max_column + 1):
        cell = current_sheet.cell(row=row_num, column=col)
        if not cell.value or not isinstance(cell.value, str):
            continue
        value = strip_if_line_marker(str(cell.value))
        if clear_all:
            for marker in (
                "{id_method}",
                "{result}",
                "{measurement_error}",
                "{measurement_method}",
                "{unit}",
                "{name_method}",
                "{group_name}",
                "{fractional_data}",
                "{fractional_unit}",
                "{fractional_result}",
                "{fractional_error}",
                "{fractional_measurement_method}",
            ):
                value = value.replace(marker, "")
            cell.value = value.strip() or None
            continue

        if "{id_method}" in value:
            value = value.replace(
                "{id_method}", method_id if method_id is not None else ""
            )
        if calc is not None:
            if "{name_method}" in value:
                value = value.replace("{name_method}", calc.research_method.name or "")
            if "{group_name}" in value:
                value = value.replace(
                    "{group_name}", _primary_group_name(calc.research_method)
                )
            if "{unit}" in value:
                value = value.replace("{unit}", calc.unit or "-")
            if "{result}" in value:
                value = value.replace(
                    "{result}", format_protocol_calculation_result(calc)
                )
            if "{measurement_error}" in value:
                value = value.replace(
                    "{measurement_error}",
                    format_measurement_error_value(calc.measurement_error),
                )
            if "{measurement_method}" in value:
                measurement_method = calc.research_method.measurement_method or "-"
                value = value.replace("{measurement_method}", measurement_method)
                adjust_cell_height_if_needed(
                    current_sheet, row_num, col, measurement_method
                )
        else:
            for marker in (
                "{result}",
                "{measurement_error}",
                "{measurement_method}",
                "{unit}",
                "{name_method}",
                "{group_name}",
            ):
                value = value.replace(marker, "")
        cell.value = value.strip() if value.strip() else None


def _hide_template_row(
    template_sheet,
    current_sheet,
    template_row: int,
    target_row: int,
    merged_cells_map,
) -> int:
    """Копирует строку шаблона и скрывает её (условие if line не выполнено)."""
    copy_row_formatting(
        template_sheet,
        current_sheet,
        template_row,
        target_row,
        merged_cells_map,
    )
    _fill_method_row_placeholders(
        current_sheet, target_row, method_id=None, calc=None, clear_all=True
    )
    if target_row not in current_sheet.row_dimensions:
        current_sheet.row_dimensions[target_row] = (
            openpyxl.worksheet.dimensions.RowDimension(current_sheet, target_row)
        )
    current_sheet.row_dimensions[target_row].hidden = True
    return target_row + 1


def _fill_fractional_placeholders(
    current_sheet,
    row_num: int,
    *,
    method_id: str | None,
    method_name: str,
    group_name: str,
    data_name: str,
    unit: str,
    result: str,
    error: str,
    measurement_method: str,
) -> None:
    """Подставляет fractional-метки и {name_method}/{group_name} как есть."""
    for col in range(1, current_sheet.max_column + 1):
        cell = current_sheet.cell(row=row_num, column=col)
        if not cell.value or not isinstance(cell.value, str):
            continue
        value = strip_if_line_marker(str(cell.value))
        if "{id_method}" in value:
            value = value.replace(
                "{id_method}", method_id if method_id is not None else ""
            )
        value = value.replace("{fractional_data}", data_name)
        value = value.replace("{fractional_unit}", unit)
        value = value.replace("{fractional_result}", result)
        value = value.replace("{fractional_error}", error)
        if "{name_method}" in value:
            value = value.replace("{name_method}", method_name)
        if "{group_name}" in value:
            value = value.replace("{group_name}", group_name)
        if "{unit}" in value:
            value = value.replace("{unit}", unit)
        if "{fractional_measurement_method}" in value:
            value = value.replace("{fractional_measurement_method}", measurement_method)
            adjust_cell_height_if_needed(
                current_sheet, row_num, col, measurement_method
            )
        for marker in (
            "{result}",
            "{measurement_error}",
            "{measurement_method}",
        ):
            value = value.replace(marker, "")
        cell.value = value.strip() if value.strip() else None


def _iter_fractional_oil_rows(calc: Calculation) -> list[tuple[str, str, str, str]]:
    """Строки фракционного состава нефти: (name, unit, result, error)."""
    try:
        if hasattr(calc, "intermediate_data") and calc.intermediate_data:
            if isinstance(calc.intermediate_data, str):
                result_data = orjson.loads(calc.intermediate_data)
            else:
                result_data = calc.intermediate_data
        else:
            if isinstance(calc.result, str):
                result_data = orjson.loads(calc.result)
            else:
                result_data = calc.result

        if isinstance(result_data, dict) and "_fractional_data" in result_data:
            fractional_data = result_data["_fractional_data"]
            combined_data = {}
            for card_data in fractional_data.values():
                if isinstance(card_data, dict):
                    for field, value in card_data.items():
                        if (
                            field not in combined_data
                            or combined_data[field] is None
                            or combined_data[field] == ""
                        ):
                            combined_data[field] = value
            result_data = combined_data

        if isinstance(result_data, dict):
            normalized_data = {}
            for key, value in result_data.items():
                normalized_key = key.replace("Температура н,к.", "Температура н.к.")
                normalized_data[normalized_key] = value
            result_data = normalized_data
    except (orjson.JSONDecodeError, TypeError) as e:
        logger.error(
            f"Не удалось распарсить результат для фракционного состава нефти: {str(e)}"
        )
        return []

    if not isinstance(result_data, dict):
        return []

    error_map = {
        "Температура н.к.": "±5",
        "Выход фракций до 100 ℃": "±1,4",
        "Выход фракций до 150 ℃": "±1,4",
        "Выход фракций до 200 ℃": "±1,4",
        "Выход фракций до 250 ℃": "±1,4",
        "Выход фракций до 270 ℃": "±1,4",
        "Выход фракций до 300 ℃": "±1,4",
    }
    fields = [
        "Температура н.к.",
        "Выход фракций до 100 ℃",
        "Выход фракций до 150 ℃",
        "Выход фракций до 200 ℃",
        "Выход фракций до 250 ℃",
        "Выход фракций до 270 ℃",
        "Выход фракций до 300 ℃",
    ]
    rows: list[tuple[str, str, str, str]] = []
    rows.append(("Фракционный состав:", "", "", ""))
    output_header_done = False
    for field in fields:
        raw = result_data.get(field)
        if raw is None or raw == "" or raw == "-":
            continue
        if field.startswith("Выход фракций") and not output_header_done:
            rows.append(("Выход фракций до температуры:", "%", "", ""))
            output_header_done = True
        try:
            num = float(str(raw).replace(",", "."))
            if field == "Температура н.к." and num > 360:
                result_text = "выше 360"
            else:
                result_text = str(raw).replace(".", ",")
        except (ValueError, TypeError):
            result_text = str(raw).replace(".", ",")
        if field == "Температура н.к.":
            name, unit = "Температура н.к.", "°C"
        else:
            name = field.replace("Выход фракций до ", "")
            unit = ""
        rows.append((name, unit, result_text, error_map.get(field, "-")))
    return rows


def _iter_fractional_condensate_rows(
    calc: Calculation,
) -> list[tuple[str, str, str, str]]:
    """Строки фракционного состава конденсата: (name, unit, result, error)."""
    try:
        if isinstance(calc.result, str):
            result_data = orjson.loads(calc.result)
        else:
            result_data = calc.result
    except (orjson.JSONDecodeError, TypeError) as e:
        logger.error(
            f"Не удалось распарсить результат для фракционного состава конденсата: {str(e)}"
        )
        return []

    if not isinstance(result_data, dict):
        return []

    normalized = {}
    for key, value in result_data.items():
        key_l = str(key).lower()
        if key_l.startswith("температура н"):
            normalized["Температура н.к."] = value
        else:
            normalized[key] = value
    result_data = normalized

    error_map = {
        "Температура н.к.": "±5",
        "10% отгона при температуре": "±4",
        "50% отгона при температуре": "±2",
        "90% отгона при температуре": "±5",
        "Температура к.к.": None,
        "Объемная доля остатка": "±0,3",
    }
    fields = [
        "Температура н.к.",
        "5% отгона при температуре",
        "10% отгона при температуре",
        "20% отгона при температуре",
        "30% отгона при температуре",
        "40% отгона при температуре",
        "50% отгона при температуре",
        "60% отгона при температуре",
        "70% отгона при температуре",
        "80% отгона при температуре",
        "90% отгона при температуре",
        "95% отгона при температуре",
        "98% отгона при температуре",
        "Температура к.к.",
        "Объемная доля отгона",
        "Объемная доля остатка",
        "Объемная доля потерь",
    ]
    rows: list[tuple[str, str, str, str]] = []
    rows.append(("Фракционный состав:", "", "", ""))
    first_percent_done = False
    for field in fields:
        raw = result_data.get(field)
        if raw is None or raw == "" or raw == "-":
            continue
        raw_s = str(raw).replace(".", ",")
        try:
            num = float(str(raw).replace(",", "."))
            if (
                field in ("Температура н.к.", "Температура к.к.")
                or "отгона при температуре" in field
            ) and num > 360:
                result_text = "выше 360"
            else:
                result_text = raw_s
        except (ValueError, TypeError):
            result_text = raw_s
            num = None

        if "отгона при температуре" in field:
            if not first_percent_done:
                name = field[0].upper() + field[1:] if field else field
                first_percent_done = True
            else:
                name = field.split("%")[0] + "%"
            unit = "°C"
        elif field in ("Температура н.к.", "Температура к.к."):
            name, unit = field, "°C"
        elif "доля" in field.lower():
            name = field[0].upper() + field[1:] if field else field
            unit = "%"
        else:
            name = field[0].upper() + field[1:] if field else field
            unit = ""

        if field == "Температура к.к.":
            error = _condensate_kk_measurement_error(raw)
        else:
            error = error_map.get(field, "-")
            if error is None:
                error = "-"
        rows.append((name, unit, result_text, error))
    return rows


def _clear_row_horizontal_borders(
    sheet,
    row_num: int,
    *,
    clear_top: bool = False,
    clear_bottom: bool = False,
) -> None:
    """Убирает верхнюю и/или нижнюю границу у ячеек строки."""
    for col in range(1, sheet.max_column + 1):
        cell = sheet.cell(row=row_num, column=col)
        if not cell.border:
            continue
        new_border = copy(cell.border)
        if clear_top:
            new_border.top = None
        if clear_bottom:
            new_border.bottom = None
        cell.border = new_border


def _restore_row_bottom_border_from_template(
    template_sheet,
    current_sheet,
    template_row: int,
    target_row: int,
) -> None:
    """Восстанавливает нижнюю границу строки по образцу из шаблона."""
    _apply_row_bottom_from_template_edge(
        template_sheet,
        current_sheet,
        template_row,
        target_row,
        source_edge="bottom",
    )


def _template_row_has_edge_border(
    template_sheet,
    row_num: int,
    *,
    edge: str,
) -> bool:
    """Есть ли у строки шаблона горизонтальная граница (top/bottom) хоть у одной ячейки."""
    for col in range(1, template_sheet.max_column + 1):
        border = template_sheet.cell(row=row_num, column=col).border
        if not border:
            continue
        side = getattr(border, edge, None)
        if side is not None and side.style:
            return True
    return False


def _apply_row_bottom_from_template_edge(
    template_sheet,
    current_sheet,
    template_row: int,
    target_row: int,
    *,
    source_edge: str,
) -> None:
    """Копирует top/bottom границы строки шаблона как нижнюю границу целевой строки."""
    for col in range(1, template_sheet.max_column + 1):
        source = template_sheet.cell(row=template_row, column=col)
        target = current_sheet.cell(row=target_row, column=col)
        src_border = source.border
        if not src_border:
            continue
        side = getattr(src_border, source_edge, None)
        if side is None or not side.style:
            continue
        if target.border:
            new_border = copy(target.border)
        else:
            new_border = Border()
        new_border.bottom = copy(side)
        target.border = new_border


def _close_previous_visible_row_before_id_method(
    template_sheet,
    current_sheet,
    *,
    next_template_row: int,
    last_visible_output_row: int | None,
    hidden_template_rows: list[int],
) -> None:
    """
    Перед видимой строкой с {id_method} закрывает предыдущий видимый блок нижней границей.

    Сначала берёт верхнюю границу этой id-строки из шаблона (если есть),
    иначе — нижнюю границу последней скрытой строки между блоками.
    """
    if last_visible_output_row is None:
        return

    if _template_row_has_edge_border(template_sheet, next_template_row, edge="top"):
        _apply_row_bottom_from_template_edge(
            template_sheet,
            current_sheet,
            next_template_row,
            last_visible_output_row,
            source_edge="top",
        )
        return

    for hidden_row in reversed(hidden_template_rows):
        if _template_row_has_edge_border(template_sheet, hidden_row, edge="bottom"):
            _apply_row_bottom_from_template_edge(
                template_sheet,
                current_sheet,
                hidden_row,
                last_visible_output_row,
                source_edge="bottom",
            )
            return


def _expand_fractional_if_line_row(
    calc: Calculation,
    template_sheet,
    current_sheet,
    template_row: int,
    current_row: int,
    merged_cells_map,
    next_id: int,
) -> tuple[int, int]:
    """Разворачивает fractional-строку шаблона в несколько видимых строк."""
    method_name = (calc.research_method.name or "").lower()
    if "конденсат" in method_name:
        frac_rows = _iter_fractional_condensate_rows(calc)
    else:
        frac_rows = _iter_fractional_oil_rows(calc)

    if not frac_rows:
        return (
            _hide_template_row(
                template_sheet,
                current_sheet,
                template_row,
                current_row,
                merged_cells_map,
            ),
            next_id,
        )

    measurement_method = calc.research_method.measurement_method or "-"
    method_name = calc.research_method.name or ""
    group_name = _primary_group_name(calc.research_method)
    written_rows: list[int] = []
    for index, (data_name, unit, result, error) in enumerate(frac_rows):
        copy_row_formatting(
            template_sheet,
            current_sheet,
            template_row,
            current_row,
            merged_cells_map,
        )
        is_first = index == 0
        is_last = index == len(frac_rows) - 1
        method_id = str(next_id) if is_first else ""
        if is_first:
            next_id += 1
        _fill_fractional_placeholders(
            current_sheet,
            current_row,
            method_id=method_id,
            method_name=method_name,
            group_name=group_name,
            data_name=data_name,
            unit=unit,
            result=result,
            error=error,
            measurement_method=measurement_method if method_id else "",
        )
        # Внутри блока фракционного состава убираем лишние горизонтальные линии.
        if is_first:
            _clear_row_horizontal_borders(current_sheet, current_row, clear_bottom=True)
        else:
            _clear_row_horizontal_borders(
                current_sheet,
                current_row,
                clear_top=True,
                clear_bottom=not is_last,
            )
        written_rows.append(current_row)
        current_row += 1

    if written_rows:
        _restore_row_bottom_border_from_template(
            template_sheet,
            current_sheet,
            template_row,
            written_rows[-1],
        )
    return current_row, next_id


def _collect_valid_calculations(samples: List[Sample]) -> List[Calculation]:
    """Собирает расчёты для таблицы методов с фильтрацией по объекту испытаний."""
    test_objects = [sample.test_object for sample in samples if sample.test_object]
    calculations: List[Calculation] = []
    for sample in samples:
        for calc in sample.calculations:
            if calc.deleted_at is None and calc.research_method:
                calculations.append(calc)

    valid_calculations: List[Calculation] = []
    for calc in calculations:
        method_name = calc.research_method.name.lower()
        if "фракционный состав" in method_name and method_name not in [
            "фракционный состав (конденсат)",
            "фракционный состав (нефть)",
        ]:
            continue
        if check_method_name(calc.research_method.name, test_objects):
            valid_calculations.append(calc)

    def get_sort_key(calc):
        method = calc.research_method
        method_name = method.name or ""
        method_sort_order = method.sort_order if method.sort_order is not None else 0
        if method.groups and len(method.groups) > 0:
            group = method.groups[0]
            group_sort_order = (
                group.sort_order if group.sort_order is not None else method_sort_order
            )
            return (group_sort_order, method_sort_order, method_name)
        return (method_sort_order, 0, method_name)

    valid_calculations.sort(key=get_sort_key)
    return valid_calculations


def _condensate_kk_measurement_error(field_value) -> str:
    """Погрешность для Температуры к.к. фракционного состава конденсата."""
    if field_value is None:
        return "±7"
    text = str(field_value).strip().lower()
    if "выше 360" in text:
        return "-"
    try:
        numeric_value = float(str(field_value).replace(",", "."))
        if isinstance(numeric_value, (int, float)) and numeric_value > 360:
            return "-"
    except (ValueError, TypeError):
        pass
    return "±7"


async def process_methods_table(
    samples: List[Sample],
    template_sheet,
    new_sheet,
    table_start,
    merged_cells_map,
    current_row,
    selection_conditions_templates: Optional[List[Dict[str, Any]]] = None,
):
    """
    Обрабатывает таблицу 1: строки с {if line} либо заполняются
    и остаются видимыми, либо копируются скрытыми.
    """
    del selection_conditions_templates  # метки таблицы 1 не используют условия отбора
    current_sheet = new_sheet
    valid_calculations = _collect_valid_calculations(samples)
    template_last_row, _ = get_template_content_bounds(template_sheet)

    table_data_start = None
    table_data_end = None
    for row_num in range(table_start, template_last_row + 1):
        row = list(
            template_sheet.iter_rows(min_row=row_num, max_row=row_num, values_only=True)
        )[0]
        cell_values = [str(cell).strip() for cell in row if cell]
        if table_data_start is None and any(
            value == "{start_table1}" for value in cell_values
        ):
            table_data_start = row_num
            continue
        if table_data_start is not None and any(
            value in TABLE1_END_MARKERS for value in cell_values
        ):
            table_data_end = row_num
            break

    if table_data_start is None:
        return current_sheet
    if table_data_end is None:
        table_data_end = template_last_row + 1

    next_id = 1
    last_visible_output_row: int | None = None
    hidden_since_last_visible: list[int] = []

    for template_row in range(table_data_start + 1, table_data_end):
        row_values = [
            template_sheet.cell(row=template_row, column=col).value
            for col in range(1, template_sheet.max_column + 1)
        ]
        row_text = " ".join(str(v) for v in row_values if v)

        conditions = None
        for value in row_values:
            if value and isinstance(value, str):
                conditions = parse_if_line_condition(value)
                if conditions is not None:
                    break

        if conditions is None:
            copy_row_formatting(
                template_sheet,
                current_sheet,
                template_row,
                current_row,
                merged_cells_map,
            )
            last_visible_output_row = current_row
            hidden_since_last_visible = []
            current_row += 1
            continue

        calc = _find_calculation_for_if_line(conditions, valid_calculations)
        if calc is None:
            hidden_since_last_visible.append(template_row)
            current_row = _hide_template_row(
                template_sheet,
                current_sheet,
                template_row,
                current_row,
                merged_cells_map,
            )
            continue

        has_id_marker = "{id_method}" in row_text
        if has_id_marker:
            _close_previous_visible_row_before_id_method(
                template_sheet,
                current_sheet,
                next_template_row=template_row,
                last_visible_output_row=last_visible_output_row,
                hidden_template_rows=hidden_since_last_visible,
            )

        if _row_has_fractional_markers(template_sheet, template_row):
            start_output_row = current_row
            current_row, next_id = _expand_fractional_if_line_row(
                calc,
                template_sheet,
                current_sheet,
                template_row,
                current_row,
                merged_cells_map,
                next_id,
            )
            if current_row > start_output_row:
                last_visible_output_row = current_row - 1
                hidden_since_last_visible = []
            continue

        copy_row_formatting(
            template_sheet,
            current_sheet,
            template_row,
            current_row,
            merged_cells_map,
        )
        method_id = str(next_id) if has_id_marker else None
        if has_id_marker:
            next_id += 1
        _fill_method_row_placeholders(
            current_sheet,
            current_row,
            method_id=method_id,
            calc=calc,
        )
        last_visible_output_row = current_row
        hidden_since_last_visible = []
        current_row += 1

    return current_sheet


def _collect_equipment_ids_from_samples(samples: List[Sample]) -> set[int]:
    """Собирает уникальные ID оборудования из поля calculation.equipment_data для всех проб."""
    equipment_ids: set[int] = set()

    for sample in samples:
        for calc in getattr(sample, "calculations", []) or []:
            if calc.deleted_at is not None:
                continue
            data = getattr(calc, "equipment_data", None)
            if not data:
                continue

            if isinstance(data, list):
                for item in data:
                    if isinstance(item, dict) and "id" in item:
                        try:
                            equipment_ids.add(int(item["id"]))
                        except (TypeError, ValueError):
                            continue
                    elif isinstance(item, (int, str)):
                        try:
                            equipment_ids.add(int(item))
                        except (TypeError, ValueError):
                            continue
            elif isinstance(data, dict) and "id" in data:
                try:
                    equipment_ids.add(int(data["id"]))
                except (TypeError, ValueError):
                    continue
            elif isinstance(data, (int, str)):
                try:
                    equipment_ids.add(int(data))
                except (TypeError, ValueError):
                    continue

    return equipment_ids


def process_equipment_table(
    equipment_list: List[Equipment],
    template_sheet,
    current_sheet,
    table_start,
    merged_cells_map,
    current_row,
):
    """Обрабатывает таблицу с оборудованием."""
    seen_equipment = set()
    unique_equipment = []

    for equipment in equipment_list:
        ver_date = None
        if equipment.verification_date:
            if hasattr(equipment.verification_date, "date"):
                ver_date = equipment.verification_date.date().isoformat()
            else:
                ver_date = str(equipment.verification_date)

        ver_end_date = None
        if equipment.verification_end_date:
            if hasattr(equipment.verification_end_date, "date"):
                ver_end_date = equipment.verification_end_date.date().isoformat()
            else:
                ver_end_date = str(equipment.verification_end_date)

        equipment_key = (
            equipment.name or "",
            equipment.serial_number or "",
            equipment.verification_info or "",
            ver_date,
            ver_end_date,
        )

        if equipment_key not in seen_equipment:
            seen_equipment.add(equipment_key)
            unique_equipment.append(equipment)

    # Сортируем список оборудования по наименованию и версии
    equipment_list = sorted(
        unique_equipment, key=lambda x: (x.name or "", x.version or "")
    )

    if not equipment_list:
        return current_sheet

    template_last_row, _ = get_template_content_bounds(template_sheet)
    start_marker_row = None
    template_row_num = None

    for row_num in range(table_start, template_last_row + 1):
        row = list(
            template_sheet.iter_rows(min_row=row_num, max_row=row_num, values_only=True)
        )[0]
        if any(cell and str(cell).strip() == "{start_table2}" for cell in row):
            start_marker_row = row_num
            template_row_num = row_num + 1
            break

    if start_marker_row is None or not template_row_num:
        return current_sheet

    if template_row_num > template_last_row:
        return current_sheet
    end_check = template_sheet.cell(row=template_row_num, column=1).value
    if end_check and str(end_check).strip() == "{end_table2}":
        return current_sheet

    idx = 1
    for equipment in equipment_list:
        copy_row_formatting(
            template_sheet,
            current_sheet,
            template_row_num,
            current_row,
            merged_cells_map,
        )

        for col in range(1, template_sheet.max_column + 1):
            cell = current_sheet.cell(row=current_row, column=col)
            if not cell.value:
                continue

            value = str(cell.value)
            if "{id_equipment}" in value:
                cell.value = value.replace("{id_equipment}", str(idx))
            elif "{name_equipment}" in value:
                equipment_name = equipment.name or ""
                cell.value = value.replace("{name_equipment}", equipment_name)
                adjust_cell_height_if_needed(
                    current_sheet, current_row, col, equipment_name
                )
            elif "{serial_num}" in value:
                cell.value = value.replace(
                    "{serial_num}", equipment.serial_number or ""
                )
            elif "{ver_info}" in value:
                verification_info = equipment.verification_info or ""
                cell.value = value.replace("{ver_info}", verification_info)
                adjust_cell_height_if_needed(
                    current_sheet, current_row, col, verification_info
                )
            elif "{ver_date}" in value:
                if equipment.verification_date:
                    formatted_date = pendulum.instance(
                        equipment.verification_date
                    ).format("DD.MM.YYYY")
                    cell.value = value.replace("{ver_date}", formatted_date)
                else:
                    cell.value = value.replace("{ver_date}", "")
            elif "{ver_end_date}" in value:
                if equipment.verification_end_date:
                    formatted_date = pendulum.instance(
                        equipment.verification_end_date
                    ).format("DD.MM.YYYY")
                    cell.value = value.replace("{ver_end_date}", formatted_date)
                else:
                    cell.value = value.replace("{ver_end_date}", "")

        current_row += 1
        idx += 1

    return current_sheet


def process_nd_table(
    samples: List[Sample],
    template_sheet,
    current_sheet,
    table_start,
    merged_cells_map,
    current_row,
):
    """Обрабатывает таблицу с нормативными документами."""
    calculations = []
    for sample in samples:
        for calc in sample.calculations:
            if calc.deleted_at is None and calc.research_method:
                calculations.append(calc)

    test_objects = []
    for sample in samples:
        if sample.test_object:
            test_objects.append(sample.test_object)

    valid_calculations = [
        calc
        for calc in calculations
        if check_method_name(calc.research_method.name, test_objects)
    ]

    def get_sort_key(calc):
        method = calc.research_method
        method_name = method.name or ""
        method_sort_order = method.sort_order if method.sort_order is not None else 0
        if method.groups and len(method.groups) > 0:
            group = method.groups[0]
            group_sort_order = (
                group.sort_order if group.sort_order is not None else method_sort_order
            )
            return (group_sort_order, method_sort_order, method_name)
        return (method_sort_order, 0, method_name)

    valid_calculations.sort(key=get_sort_key)

    nd_list = []
    seen_nd = set()
    for calc in valid_calculations:
        nd_code = calc.research_method.nd_code or ""
        nd_name = calc.research_method.nd_name or ""
        nd_key = (nd_code, nd_name)
        if nd_key not in seen_nd and (nd_code or nd_name):
            seen_nd.add(nd_key)
            nd_list.append(nd_key)

    if not nd_list:
        return current_sheet

    template_last_row, _ = get_template_content_bounds(template_sheet)
    start_marker_row = None
    template_row_num = None

    for row_num in range(table_start, template_last_row + 1):
        row = list(
            template_sheet.iter_rows(min_row=row_num, max_row=row_num, values_only=True)
        )[0]
        if any(cell and str(cell).strip() == "{start_table3}" for cell in row):
            start_marker_row = row_num
            template_row_num = row_num + 1
            break

    if start_marker_row is None or not template_row_num:
        return current_sheet

    if template_row_num > template_last_row:
        return current_sheet
    end_check = template_sheet.cell(row=template_row_num, column=1).value
    if end_check and str(end_check).strip() == "{end_table3}":
        return current_sheet

    idx = 1
    for nd_code, nd_name in nd_list:
        copy_row_formatting(
            template_sheet,
            current_sheet,
            template_row_num,
            current_row,
            merged_cells_map,
        )

        for col in range(1, template_sheet.max_column + 1):
            cell = current_sheet.cell(row=current_row, column=col)
            if not cell.value:
                continue

            value = str(cell.value)
            if "{id_nd}" in value:
                cell.value = value.replace("{id_nd}", str(idx))
            elif "{nd_code}" in value:
                cell.value = value.replace("{nd_code}", nd_code)
            elif "{name_nd}" in value:
                cell.value = value.replace("{name_nd}", nd_name)
                adjust_cell_height_if_needed(current_sheet, current_row, col, nd_name)

        current_row += 1
        idx += 1

    return current_sheet


async def generate_protocol_excel(db: AsyncSession, protocol_id: int) -> Response:
    """Генерирует Excel файл протокола."""
    try:
        query = (
            select(Protocol)
            .where(Protocol.id == protocol_id)
            .options(
                selectinload(Protocol.laboratory),
                selectinload(Protocol.department),
                selectinload(Protocol.protocol_template),
            )
        )
        result = await db.execute(query)
        protocol = result.scalar_one_or_none()

        if not protocol:
            raise NotFoundError(f"Протокол с id {protocol_id} не найден")

        if not protocol.protocol_template:
            raise ValidationError("У протокола отсутствует шаблон Excel")

        sample_ids = protocol.samples or []
        if not sample_ids:
            raise ValidationError("У протокола отсутствуют пробы")

        samples_query = (
            select(Sample)
            .where(Sample.id.in_(sample_ids))
            .where(Sample.deleted_at.is_(None))
            .options(
                selectinload(Sample.branch),
                selectinload(Sample.sampling_location),
                selectinload(Sample.calculations)
                .selectinload(Calculation.research_method)
                .selectinload(ResearchMethod.groups),
            )
        )
        samples_result = await db.execute(samples_query)
        samples = list(samples_result.scalars().all())

        if not samples:
            raise ValidationError("У протокола отсутствуют пробы")

        # Загружаем шаблоны условий отбора для лаборатории/подразделения
        selection_conditions_templates = []
        if protocol.laboratory_id or protocol.department_id:
            selection_conditions_query = select(SelectionConditions).where(
                SelectionConditions.deleted_at.is_(None)
            )
            if protocol.laboratory_id:
                selection_conditions_query = selection_conditions_query.where(
                    SelectionConditions.laboratory_id == protocol.laboratory_id
                )
            if protocol.department_id:
                selection_conditions_query = selection_conditions_query.where(
                    SelectionConditions.department_id == protocol.department_id
                )
            selection_conditions_result = await db.execute(selection_conditions_query)
            selection_conditions_list = list(
                selection_conditions_result.scalars().all()
            )
            if selection_conditions_list:
                selection_conditions_templates = [
                    {"conditions": sc.conditions} for sc in selection_conditions_list
                ]

        # Собираем список используемого оборудования для всех расчетов по пробам
        equipment_ids = _collect_equipment_ids_from_samples(samples)
        equipment_list: List[Equipment] = []
        if equipment_ids:
            equipment_query = select(Equipment).where(Equipment.id.in_(equipment_ids))
            equipment_result = await db.execute(equipment_query)
            equipment_list = list(equipment_result.scalars().all())

        abbreviations_by_name = await get_protocol_abbreviations_by_names(
            db,
            [sample.test_object for sample in samples if sample.test_object],
        )
        protocol_abbreviation = pick_first_protocol_abbreviation(
            abbreviations_by_name,
            [sample.test_object for sample in samples],
        )
        _protocol_abbreviation_ctx.set(protocol_abbreviation)

        file_data = protocol.protocol_template.file
        try:
            template_bytes = BytesIO(base64.b64decode(file_data))
        except Exception:
            try:
                with open(file_data, "rb") as f:
                    template_bytes = BytesIO(f.read())
            except Exception:
                template_bytes = BytesIO(
                    file_data.encode() if isinstance(file_data, str) else file_data
                )

        template_workbook = openpyxl.load_workbook(template_bytes)
        template_sheet = template_workbook.active
        sampling_location_name_only = template_contains_marker(template_sheet, "{mode}")
        template_last_row, template_last_col = get_template_content_bounds(
            template_sheet
        )

        new_workbook = openpyxl.Workbook()
        new_sheet = new_workbook.active
        new_sheet.title = "Лист 1"
        copy_sheet_page_settings(template_sheet, new_sheet)

        if hasattr(template_sheet, "oddFooter") and hasattr(new_sheet, "oddFooter"):
            new_sheet.oddFooter.left = await process_footer_test_protocol_number(
                protocol, samples, template_sheet.oddFooter.left
            )
            new_sheet.oddFooter.center = await process_footer_test_protocol_number(
                protocol, samples, template_sheet.oddFooter.center
            )
            new_sheet.oddFooter.right = await process_footer_test_protocol_number(
                protocol, samples, template_sheet.oddFooter.right
            )
        if hasattr(template_sheet, "evenFooter") and hasattr(new_sheet, "evenFooter"):
            new_sheet.evenFooter.left = await process_footer_test_protocol_number(
                protocol, samples, template_sheet.evenFooter.left
            )
            new_sheet.evenFooter.center = await process_footer_test_protocol_number(
                protocol, samples, template_sheet.evenFooter.center
            )
            new_sheet.evenFooter.right = await process_footer_test_protocol_number(
                protocol, samples, template_sheet.evenFooter.right
            )
        if hasattr(template_sheet, "firstFooter") and hasattr(new_sheet, "firstFooter"):
            new_sheet.firstFooter.left = await process_footer_test_protocol_number(
                protocol, samples, template_sheet.firstFooter.left
            )
            new_sheet.firstFooter.center = await process_footer_test_protocol_number(
                protocol, samples, template_sheet.firstFooter.center
            )
            new_sheet.firstFooter.right = await process_footer_test_protocol_number(
                protocol, samples, template_sheet.firstFooter.right
            )

        merged_cells_map = new_sheet.merged_cells

        header_end = await process_header(
            protocol,
            samples,
            template_sheet,
            new_sheet,
            merged_cells_map,
            selection_conditions_templates,
            sampling_location_name_only,
        )
        if header_end == 0:
            raise ValidationError("Ошибка при обработке шапки протокола")

        table_start = await process_header_and_conditions(
            protocol,
            samples,
            template_sheet,
            new_sheet,
            header_end,
            merged_cells_map,
            selection_conditions_templates,
            sampling_location_name_only,
        )

        current_sheet = await process_methods_table(
            samples,
            template_sheet,
            new_sheet,
            table_start,
            merged_cells_map,
            new_sheet.max_row + 1,
            selection_conditions_templates,
        )
        if not current_sheet:
            raise ValidationError("Ошибка при обработке таблицы методов")

        table1_end = None
        for row_num in range(table_start, template_last_row + 1):
            row = list(
                template_sheet.iter_rows(
                    min_row=row_num, max_row=row_num, values_only=True
                )
            )[0]
            if any(cell and str(cell).strip() in TABLE1_END_MARKERS for cell in row):
                table1_end = row_num
                break

        if table1_end:
            current_sheet, current_row = await process_between_tables(
                protocol,
                samples,
                template_sheet,
                current_sheet,
                table1_end,
                merged_cells_map,
                current_sheet.max_row + 1,
                selection_conditions_templates,
                sampling_location_name_only,
            )

            current_sheet = process_equipment_table(
                equipment_list,
                template_sheet,
                current_sheet,
                table1_end + 1,
                merged_cells_map,
                current_row,
            )
            if not current_sheet:
                raise ValidationError("Ошибка при обработке таблицы оборудования")

            table2_end = None
            for row_num in range(table1_end + 1, template_last_row + 1):
                row = list(
                    template_sheet.iter_rows(
                        min_row=row_num, max_row=row_num, values_only=True
                    )
                )[0]
                if any(cell and str(cell).strip() == "{end_table2}" for cell in row):
                    table2_end = row_num
                    break

            if table2_end:
                current_sheet, current_row = await process_between_tables(
                    protocol,
                    samples,
                    template_sheet,
                    current_sheet,
                    table2_end,
                    merged_cells_map,
                    current_sheet.max_row + 1,
                    selection_conditions_templates,
                    sampling_location_name_only,
                )

                current_sheet = process_nd_table(
                    samples,
                    template_sheet,
                    current_sheet,
                    table2_end + 1,
                    merged_cells_map,
                    current_row,
                )
                if not current_sheet:
                    raise ValidationError("Ошибка при обработке таблицы НД")

                table3_end = None
                for row_num in range(table2_end + 1, template_last_row + 1):
                    row = list(
                        template_sheet.iter_rows(
                            min_row=row_num, max_row=row_num, values_only=True
                        )
                    )[0]
                    if any(
                        cell and str(cell).strip() == "{end_table3}" for cell in row
                    ):
                        table3_end = row_num
                        break

                if table3_end:
                    current_sheet = await process_footer(
                        protocol,
                        samples,
                        template_sheet,
                        current_sheet,
                        table3_end,
                        merged_cells_map,
                        current_sheet.max_row + 1,
                        selection_conditions_templates,
                        sampling_location_name_only,
                    )

        copy_sheet_page_settings(template_sheet, new_sheet)
        copy_column_dimensions(template_sheet, new_sheet)

        apply_sheet_print_area(new_sheet)

        output = BytesIO()
        new_workbook.save(output)
        output.seek(0)

        # Формируем имя файла: "Протокол_<Номер протокола>.xlsx"
        protocol_number = None
        if protocol.test_protocol_number:
            protocol_number_str = str(protocol.test_protocol_number).strip()
            if protocol_number_str:
                protocol_number = protocol_number_str

        # Если номер протокола отсутствует, используем ID протокола
        if not protocol_number:
            protocol_number = str(protocol_id)
            logger.warning(
                f"Номер протокола отсутствует для протокола {protocol_id}, используется ID"
            )

        clean_protocol_number = re.sub(r'[<>:"/\\|?*\x00-\x1F]', "_", protocol_number)
        clean_protocol_number = clean_protocol_number.replace(" ", "_")

        # Формируем имя файла с русскими символами
        filename_ru = f"Протокол_{clean_protocol_number}.xlsx"

        # Кодируем русское имя файла для HTTP‑заголовка (RFC 5987)
        encoded_filename = quote(filename_ru, safe="", encoding="utf-8")

        # Используем только filename* для избежания проблем с latin-1 кодированием в Starlette
        content_disposition = f"attachment; filename*=UTF-8''{encoded_filename}"

        return Response(
            content=output.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": content_disposition},
        )

    except NotFoundError as e:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=str(e))
    except ValidationError as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(e))
    except Exception as e:
        logger.error(f"Ошибка при генерации протокола: {str(e)}")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Ошибка при генерации протокола: {str(e)}",
        )
