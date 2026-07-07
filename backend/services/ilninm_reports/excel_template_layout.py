import openpyxl.worksheet.worksheet


def insert_rows_for_data_count(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    template_data_row: int,
    data_row_count: int,
) -> None:
    """
    Резервирует место под строки данных: хвост шаблона (под образцовой строкой)
    сдвигается вниз, чтобы не затираться при заполнении.
    """
    if data_row_count <= 1:
        return
    ws.insert_rows(template_data_row + 1, data_row_count - 1)


def unmerge_cells_in_row_range(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    first_row: int,
    last_row: int,
    max_col: int,
) -> None:
    """
    Снимает объединения только в блоке строк данных, не в хвосте шаблона.

    После insert_rows openpyxl иногда падает с KeyError при unmerge — игнорируем.
    """
    to_unmerge: list[str] = []
    for merged in list(ws.merged_cells.ranges):
        if merged.max_row < first_row or merged.min_row > last_row:
            continue
        if merged.min_col > max_col or merged.max_col < 1:
            continue
        to_unmerge.append(str(merged))
    for range_str in to_unmerge:
        try:
            ws.unmerge_cells(range_str)
        except KeyError:
            pass
