import base64
from io import BytesIO
from typing import Any
import openpyxl
from openpyxl.cell.cell import Cell
from openpyxl.styles import Font
from sqlalchemy.ext.asyncio import AsyncSession
from services.ilninm_reports.constants import (
    PHYSICOCHEMICAL_PLACEHOLDER_PERIOD,
    PHYSICOCHEMICAL_PLACEHOLDER_SAMPLING_LOCATION,
    PHYSICOCHEMICAL_TEMPLATE_DATA_ROW,
    PHYSICOCHEMICAL_TEMPLATE_HEADER_LAST_ROW,
    REPORT_EMPTY_CELL_VALUE,
)
from services.ilninm_reports.physicochemical import (
    METHOD_COLUMNS,
    format_report_period,
    get_physicochemical_report_rows,
    resolve_sampling_location_display,
)
from services.ilninm_reports.sample_count_generator import (
    REPORT_FONT_NAME,
    REPORT_FONT_SIZE,
)
from utils.protocol_generator_utils import (
    copy_cell_style,
    copy_column_dimensions,
    copy_row_formatting,
)


def _apply_report_font(cell: Cell) -> None:
    old = cell.font
    if old:
        cell.font = Font(
            name=REPORT_FONT_NAME,
            size=REPORT_FONT_SIZE,
            bold=old.bold,
            italic=old.italic,
            underline=old.underline,
            strike=old.strike,
            color="FF000000",
        )
        return
    cell.font = Font(name=REPORT_FONT_NAME, size=REPORT_FONT_SIZE, color="FF000000")


def _merged_cell_anchor(
    ws: openpyxl.worksheet.worksheet.Worksheet, row: int, col: int
) -> tuple[int, int]:
    """Возвращает верхнюю левую ячейку объединённого диапазона."""
    for merged in ws.merged_cells.ranges:
        if (
            merged.min_row <= row <= merged.max_row
            and merged.min_col <= col <= merged.max_col
        ):
            return merged.min_row, merged.min_col
    return row, col


def _apply_placeholders_to_cell(
    cell: Cell,
    period_text: str,
    sampling_location_text: str,
) -> None:
    if cell.value is None:
        return
    text = str(cell.value)
    updated = text
    if PHYSICOCHEMICAL_PLACEHOLDER_PERIOD in updated:
        updated = updated.replace(PHYSICOCHEMICAL_PLACEHOLDER_PERIOD, period_text)
    if PHYSICOCHEMICAL_PLACEHOLDER_SAMPLING_LOCATION in updated:
        updated = updated.replace(
            PHYSICOCHEMICAL_PLACEHOLDER_SAMPLING_LOCATION,
            sampling_location_text,
        )
    if updated != text:
        cell.value = updated


def _replace_header_placeholders(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    period_text: str,
    sampling_location_text: str,
) -> None:
    """Подставляет период и место отбора в шапку шаблона (строки до строки данных)."""
    processed: set[tuple[int, int]] = set()
    max_col = max(ws.max_column, 30)
    for row in range(1, PHYSICOCHEMICAL_TEMPLATE_HEADER_LAST_ROW + 1):
        for col in range(1, max_col + 1):
            anchor_row, anchor_col = _merged_cell_anchor(ws, row, col)
            anchor = (anchor_row, anchor_col)
            if anchor in processed:
                continue
            processed.add(anchor)
            _apply_placeholders_to_cell(
                ws.cell(row=anchor_row, column=anchor_col),
                period_text,
                sampling_location_text,
            )


async def build_physicochemical_excel(
    db: AsyncSession,
    template_file_base64: str,
    laboratory_id: int,
    sampling_date_from: Any,
    sampling_date_to: Any,
    sampling_location: str,
    department_id: int | None = None,
) -> bytes:
    """
    Строит Excel по шаблону: шапка с периодом и местом отбора,
    строки данных с 6-й строки по образцу шаблона.
    """
    template_bytes = base64.b64decode(template_file_base64)
    wb = openpyxl.load_workbook(BytesIO(template_bytes))
    ws = wb.worksheets[0] if wb.worksheets else wb.active

    period_text = format_report_period(sampling_date_from, sampling_date_to)
    location_text = resolve_sampling_location_display(sampling_location)
    for sheet in wb.worksheets:
        _replace_header_placeholders(sheet, period_text, location_text)

    rows = await get_physicochemical_report_rows(
        db,
        laboratory_id,
        department_id,
        sampling_date_from,
        sampling_date_to,
        sampling_location,
    )

    template_row = PHYSICOCHEMICAL_TEMPLATE_DATA_ROW
    max_col = max((spec.column for spec in METHOD_COLUMNS), default=19)
    output_row = template_row

    for index, row_data in enumerate(rows, start=1):
        if output_row != template_row:
            copy_row_formatting(ws, ws, template_row, output_row, merged_cells_map=None)
            for col in range(1, max_col + 1):
                src = ws.cell(row=template_row, column=col)
                tgt = ws.cell(row=output_row, column=col)
                copy_cell_style(src, tgt)

        ws.cell(row=output_row, column=1).value = index
        ws.cell(row=output_row, column=2).value = row_data.well
        ws.cell(row=output_row, column=3).value = row_data.sampling_date
        for spec in METHOD_COLUMNS:
            ws.cell(row=output_row, column=spec.column).value = (
                row_data.values_by_column.get(spec.column, REPORT_EMPTY_CELL_VALUE)
            )

        for col in range(1, max_col + 1):
            _apply_report_font(ws.cell(row=output_row, column=col))

        output_row += 1

    copy_column_dimensions(ws, ws)

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
