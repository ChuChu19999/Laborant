import base64
from io import BytesIO
from typing import Any
import openpyxl
from openpyxl.cell.cell import Cell
from openpyxl.styles import Font
from sqlalchemy.ext.asyncio import AsyncSession
from services.ilninm_reports.constants import (
    KGS_AVERAGE_ROW_LABEL,
    KGS_PLACEHOLDER_PERIOD,
    KGS_TEMPLATE_DATA_ROW,
    KGS_TEMPLATE_HEADER_LAST_ROW,
    REPORT_EMPTY_CELL_VALUE,
)
from services.ilninm_reports.excel_template_layout import (
    insert_rows_for_data_count,
    unmerge_cells_in_row_range,
)
from services.ilninm_reports.kgs import (
    KGS_METHOD_COLUMNS,
    format_report_period,
    get_kgs_report_groups,
)
from services.ilninm_reports.sample_count_generator import (
    REPORT_FONT_NAME,
    REPORT_FONT_SIZE,
)
from utils.protocol_generator_utils import (
    copy_cell_style,
    copy_column_dimensions,
    copy_row_formatting,
)


def _apply_report_font(cell: Cell) -> None:
    old = cell.font
    if old:
        cell.font = Font(
            name=REPORT_FONT_NAME,
            size=REPORT_FONT_SIZE,
            bold=old.bold,
            italic=old.italic,
            underline=old.underline,
            strike=old.strike,
            color="FF000000",
        )
        return
    cell.font = Font(name=REPORT_FONT_NAME, size=REPORT_FONT_SIZE, color="FF000000")


def _merged_cell_anchor(ws: openpyxl.worksheet.worksheet.Worksheet, row: int, col: int) -> tuple[int, int]:
    for merged in ws.merged_cells.ranges:
        if merged.min_row <= row <= merged.max_row and merged.min_col <= col <= merged.max_col:
            return merged.min_row, merged.min_col
    return row, col


def _writable_cell(ws: openpyxl.worksheet.worksheet.Worksheet, row: int, col: int) -> Cell:
    """Ячейка, в которую можно записать значение (верхняя левая при объединении)."""
    anchor_row, anchor_col = _merged_cell_anchor(ws, row, col)
    return ws.cell(row=anchor_row, column=anchor_col)


def _replace_period_placeholder(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    period_text: str,
) -> None:
    processed: set[tuple[int, int]] = set()
    max_col = max(ws.max_column, 8)
    for row in range(1, KGS_TEMPLATE_HEADER_LAST_ROW + 1):
        for col in range(1, max_col + 1):
            anchor_row, anchor_col = _merged_cell_anchor(ws, row, col)
            anchor = (anchor_row, anchor_col)
            if anchor in processed:
                continue
            processed.add(anchor)
            cell = ws.cell(row=anchor_row, column=anchor_col)
            if cell.value is None:
                continue
            text = str(cell.value)
            if KGS_PLACEHOLDER_PERIOD in text:
                cell.value = text.replace(KGS_PLACEHOLDER_PERIOD, period_text)


def _write_data_row(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    template_row: int,
    output_row: int,
    max_col: int,
    row_index: int | None,
    location_display: str,
    sampling_date: str,
    values_by_column: dict[int, str],
) -> None:
    if output_row != template_row:
        copy_row_formatting(ws, ws, template_row, output_row, merged_cells_map=None)
        for col in range(1, max_col + 1):
            copy_cell_style(
                ws.cell(row=template_row, column=col),
                ws.cell(row=output_row, column=col),
            )

    _writable_cell(ws, output_row, 1).value = row_index
    _writable_cell(ws, output_row, 2).value = location_display
    _writable_cell(ws, output_row, 3).value = sampling_date
    for spec in KGS_METHOD_COLUMNS:
        _writable_cell(ws, output_row, spec.column).value = values_by_column.get(spec.column, REPORT_EMPTY_CELL_VALUE)

    for col in range(1, max_col + 1):
        _apply_report_font(_writable_cell(ws, output_row, col))


def _write_average_row(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    template_row: int,
    output_row: int,
    max_col: int,
    values_by_column: dict[int, str],
) -> None:
    if output_row != template_row:
        copy_row_formatting(ws, ws, template_row, output_row, merged_cells_map=None)
        for col in range(1, max_col + 1):
            copy_cell_style(
                ws.cell(row=template_row, column=col),
                ws.cell(row=output_row, column=col),
            )

    _writable_cell(ws, output_row, 1).value = None
    _writable_cell(ws, output_row, 2).value = KGS_AVERAGE_ROW_LABEL
    _writable_cell(ws, output_row, 3).value = None
    for spec in KGS_METHOD_COLUMNS:
        _writable_cell(ws, output_row, spec.column).value = values_by_column.get(spec.column, REPORT_EMPTY_CELL_VALUE)

    for col in range(1, max_col + 1):
        _apply_report_font(_writable_cell(ws, output_row, col))


async def build_kgs_excel(
    db: AsyncSession,
    template_file_base64: str,
    laboratory_id: int,
    sampling_date_from: Any,
    sampling_date_to: Any,
    department_id: int | None = None,
) -> bytes:
    """Строит Excel по шаблону: шапка с периодом, строки данных с 18-й строки."""
    template_bytes = base64.b64decode(template_file_base64)
    wb = openpyxl.load_workbook(BytesIO(template_bytes))
    ws = wb.worksheets[0] if wb.worksheets else wb.active

    period_text = format_report_period(sampling_date_from, sampling_date_to)
    for sheet in wb.worksheets:
        _replace_period_placeholder(sheet, period_text)

    groups = await get_kgs_report_groups(
        db,
        laboratory_id,
        department_id,
        sampling_date_from,
        sampling_date_to,
    )

    template_row = KGS_TEMPLATE_DATA_ROW
    max_col = max((spec.column for spec in KGS_METHOD_COLUMNS), default=8)
    data_row_count = sum(len(group.data_rows) + 1 for group in groups)
    insert_rows_for_data_count(ws, template_row, data_row_count)
    if data_row_count > 0:
        unmerge_cells_in_row_range(
            ws,
            template_row,
            template_row + data_row_count - 1,
            max_col,
        )
    output_row = template_row
    row_index = 1

    for group in groups:
        for data_row in group.data_rows:
            _write_data_row(
                ws,
                template_row,
                output_row,
                max_col,
                row_index,
                data_row.location_display,
                data_row.sampling_date,
                data_row.values_by_column,
            )
            row_index += 1
            output_row += 1

        _write_average_row(
            ws,
            template_row,
            output_row,
            max_col,
            group.average_row.values_by_column,
        )
        output_row += 1

    copy_column_dimensions(ws, ws)

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
