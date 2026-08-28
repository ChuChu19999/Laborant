from __future__ import annotations
import base64
from copy import copy
from io import BytesIO
from typing import Any
import openpyxl
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Side
from openpyxl.worksheet.worksheet import Worksheet
import pendulum
from sqlalchemy.ext.asyncio import AsyncSession
from services.ilninm_reports.common import format_report_period
from services.ilninm_reports.constants import (
    ROW_TITLE_KALIBROVOCHNAYA_NEFT_UGPU,
    ROW_TITLE_TOVARNAYA_NEFT_NGDU,
    SAMPLE_COUNT_EMPTY_CELL_VALUE,
    SAMPLE_COUNT_PLACEHOLDER_KOL_VO,
    SAMPLE_COUNT_PLACEHOLDER_PERIOD,
    SAMPLE_COUNT_ROW_TITLES,
    SAMPLE_COUNT_TEMPLATE_HEADER_ROW_COUNT,
)
from services.ilninm_reports.sample_count import (
    ROW_TITLE_TO_BRANCH,
    SampleCountReportData,
    get_sample_count_report_data,
    is_sample_count_row_visible_for_branch,
    match_row_title_to_value,
    normalize_row_title_for_match,
)
from utils.excel_typing import as_writable_cell, require_worksheet, set_cell_value
from utils.ilninm_reports.excel_helpers import apply_report_font, merged_cell_anchor
from utils.protocol.template_markers import (
    copy_cell_style,
    copy_column_dimensions,
    copy_row_formatting,
)

# Высота строки с текстом в столбце C (пункты Excel).
ROW_HEIGHT_COLUMN_C_LINES = 15
ROW_HEIGHT_COLUMN_C_SINGLE_LINE_TALL = 30

_ROW_TITLES_TALL_WHEN_SINGLE_LINE = frozenset(
    {
        ROW_TITLE_TOVARNAYA_NEFT_NGDU.lower(),
        ROW_TITLE_KALIBROVOCHNAYA_NEFT_UGPU.lower(),
    }
)


def _split_column_c_lines(value: str) -> list[str]:
    """Разбить значение столбца C на строки листа."""
    if not value or value == SAMPLE_COUNT_EMPTY_CELL_VALUE:
        return [SAMPLE_COUNT_EMPTY_CELL_VALUE]
    parts = value.split("\n")
    while parts and not parts[-1].strip():
        parts.pop()
    return parts if parts else [SAMPLE_COUNT_EMPTY_CELL_VALUE]


def _row_height_for_category(category_key: str, lines_in_category: int) -> float:
    """Подобрать высоту строки категории в зависимости от числа строк текста."""
    if lines_in_category == 1 and category_key.lower() in _ROW_TITLES_TALL_WHEN_SINGLE_LINE:
        return ROW_HEIGHT_COLUMN_C_SINGLE_LINE_TALL
    return ROW_HEIGHT_COLUMN_C_LINES


def _set_fixed_row_height(ws: Worksheet, row: int, height: float) -> None:
    """Задать фиксированную высоту строки листа."""
    ws.row_dimensions[row].height = height


def _cell_border_without_horizontal_edges(
    cell: Cell,
    *,
    remove_top: bool,
    remove_bottom: bool,
) -> None:
    """Снять верхнюю и/или нижнюю границу ячейки, боковые не трогает."""
    if not remove_top and not remove_bottom:
        return
    old = cell.border or Border()
    cell.border = Border(
        left=old.left,
        right=old.right,
        top=Side(style=None) if remove_top else old.top,
        bottom=Side(style=None) if remove_bottom else old.bottom,
    )


def _apply_multiline_category_borders(
    ws: Worksheet,
    start_row: int,
    end_row: int,
) -> None:
    """Убрать горизонтальные границы между строками одной категории в столбце C."""
    line_count = end_row - start_row + 1
    if line_count <= 1:
        return
    for idx, row in enumerate(range(start_row, end_row + 1)):
        border_cell = as_writable_cell(ws.cell(row=row, column=3))
        if border_cell is None:
            continue
        _cell_border_without_horizontal_edges(
            border_cell,
            remove_top=idx > 0,
            remove_bottom=idx < line_count - 1,
        )


def _copy_worksheet_merged_ranges(
    src_ws: Worksheet,
    tgt_ws: Worksheet,
    *,
    min_row: int,
    max_row: int,
    row_offset: int = 0,
) -> None:
    """Скопировать объединения ячеек шаблона, полностью попадающие в диапазон строк."""
    for merged_range in list(src_ws.merged_cells.ranges):
        if merged_range.min_row < min_row or merged_range.max_row > max_row:
            continue
        tgt_ws.merge_cells(
            start_row=merged_range.min_row + row_offset,
            start_column=merged_range.min_col,
            end_row=merged_range.max_row + row_offset,
            end_column=merged_range.max_col,
        )


def _apply_cell_top_border_from(source: Cell, target: Cell) -> None:
    """Скопировать верхнюю границу из образцовой ячейки."""
    src_border = source.border
    if not src_border or not src_border.top or src_border.top.style is None:
        return
    old = target.border or Border()
    target.border = Border(
        left=old.left,
        right=old.right,
        top=copy(src_border.top),
        bottom=old.bottom,
    )


def _apply_branch_block_top_border(
    template_ws: Worksheet,
    new_ws: Worksheet,
    block_start_row: int,
    template_top_row: int,
) -> None:
    """Поставить верхнюю границу блока филиала как у первой строки категорий шаблона."""
    for col in (1, 2, 3):
        source_cell = as_writable_cell(template_ws.cell(row=template_top_row, column=col))
        target_cell = as_writable_cell(new_ws.cell(row=block_start_row, column=col))
        if source_cell is None or target_cell is None:
            continue
        _apply_cell_top_border_from(source_cell, target_cell)


def _replace_text_placeholders(text: str, period_text: str, total_samples: int) -> str:
    """Подставить период и общее число проб в текст ячейки."""
    updated = text
    if SAMPLE_COUNT_PLACEHOLDER_PERIOD in updated:
        updated = updated.replace(SAMPLE_COUNT_PLACEHOLDER_PERIOD, period_text)
    if SAMPLE_COUNT_PLACEHOLDER_KOL_VO in updated:
        updated = updated.replace(SAMPLE_COUNT_PLACEHOLDER_KOL_VO, str(total_samples))
    return updated


def _replace_sample_count_header_placeholders(
    ws: Worksheet,
    period_text: str,
    total_samples: int,
) -> None:
    """Подставить период и общее число проб в метки шапки шаблона."""
    processed: set[tuple[int, int]] = set()
    max_col = max(ws.max_column, 3)
    for row in range(1, SAMPLE_COUNT_TEMPLATE_HEADER_ROW_COUNT + 1):
        for col in range(1, max_col + 1):
            anchor_row, anchor_col = merged_cell_anchor(ws, row, col)
            anchor = (anchor_row, anchor_col)
            if anchor in processed:
                continue
            processed.add(anchor)
            cell = ws.cell(row=anchor_row, column=anchor_col)
            if cell.value is None:
                continue
            text = str(cell.value)
            replaced = _replace_text_placeholders(text, period_text, total_samples)
            if replaced != text:
                cell.value = replaced


def _copy_template_header_rows(
    template_ws: Worksheet,
    new_ws: Worksheet,
    period_text: str,
    total_samples: int,
) -> int:
    """Скопировать шапку шаблона на новый лист и вернуть номер следующей строки."""
    max_col = max(template_ws.max_column, 3)
    for row in range(1, SAMPLE_COUNT_TEMPLATE_HEADER_ROW_COUNT + 1):
        copy_row_formatting(template_ws, new_ws, row, row, merged_cells_map=None)
        for col in range(1, max_col + 1):
            anchor_row, anchor_col = merged_cell_anchor(template_ws, row, col)
            src = template_ws.cell(row=anchor_row, column=anchor_col)
            tgt = new_ws.cell(row=row, column=col)
            copy_cell_style(src, tgt)
            if row == anchor_row and col == anchor_col and src.value is not None:
                writable_tgt = as_writable_cell(tgt)
                if writable_tgt is not None:
                    writable_tgt.value = _replace_text_placeholders(str(src.value), period_text, total_samples)
        if template_ws.row_dimensions[row].height is not None:
            new_ws.row_dimensions[row].height = template_ws.row_dimensions[row].height
    _copy_worksheet_merged_ranges(
        template_ws,
        new_ws,
        min_row=1,
        max_row=SAMPLE_COUNT_TEMPLATE_HEADER_ROW_COUNT,
    )
    return SAMPLE_COUNT_TEMPLATE_HEADER_ROW_COUNT + 1


def _get_cell_b_value(ws: Worksheet, row: int) -> Any:
    """Вернуть значение столбца B с учётом объединённых ячеек."""
    cell = ws.cell(row=row, column=2)
    if cell.value is not None:
        return cell.value
    for merged in ws.merged_cells.ranges:
        if merged.min_col <= 2 <= merged.max_col and merged.min_row <= row <= merged.max_row:
            return ws.cell(row=merged.min_row, column=2).value
    return None


def _find_template_data_rows(
    ws: Worksheet,
) -> list[tuple[int, Any]]:
    """Найти в шаблоне строки категорий по значениям столбца B."""
    result = []
    for row_idx in range(1, ws.max_row + 1):
        val = _get_cell_b_value(ws, row_idx)
        key = normalize_row_title_for_match(val)
        if not key:
            continue
        key_lower = key.lower()
        for title in SAMPLE_COUNT_ROW_TITLES:
            if title.lower() == key_lower:
                result.append((row_idx, val))
                break
    if not result:
        return result
    min_row = result[0][0]
    max_row = result[-1][0]
    expanded = []
    for merged in ws.merged_cells.ranges:
        if merged.min_col > 2 or merged.max_col < 2:
            continue
        if merged.max_row <= min_row or merged.min_row >= max_row:
            continue
        top_val = ws.cell(row=merged.min_row, column=2).value
        if top_val is None:
            continue
        for r in range(merged.min_row, merged.max_row + 1):
            if any(t[0] == r for t in result):
                continue
            expanded.append((r, top_val))
    result = sorted(result + expanded, key=lambda x: x[0])
    return result


async def build_sample_count_excel(
    db: AsyncSession,
    template_file_base64: str,
    laboratory_id: int,
    receiving_date_from: pendulum.DateTime | None,
    receiving_date_to: pendulum.DateTime | None,
    department_id: int | None = None,
) -> tuple[bytes, bytes]:
    """Загрузить данные отчёта «Количество проб» и собрать Excel + диагностику."""
    report_data = await get_sample_count_report_data(
        db,
        laboratory_id=laboratory_id,
        receiving_date_from=receiving_date_from,
        receiving_date_to=receiving_date_to,
        department_id=department_id,
    )
    return render_sample_count_excel(
        template_file_base64,
        receiving_date_from,
        receiving_date_to,
        report_data,
    )


def render_sample_count_excel(
    template_file_base64: str,
    receiving_date_from: pendulum.DateTime | None,
    receiving_date_to: pendulum.DateTime | None,
    report_data: SampleCountReportData,
) -> tuple[bytes, bytes]:
    """Собрать байты Excel-отчёта и текст диагностики по готовым данным."""
    template_bytes = BytesIO(base64.b64decode(template_file_base64))
    template_wb = openpyxl.load_workbook(template_bytes)
    template_ws = require_worksheet(template_wb)
    data_rows = _find_template_data_rows(template_ws)
    if not data_rows:
        empty_txt = "Диагностика не сформирована: в шаблоне не найдены строки категорий.\n"
        return template_bytes.getvalue(), empty_txt.encode("utf-8")

    by_branch = report_data.by_branch
    total_samples = report_data.total_samples
    diagnostics_txt = (report_data.diagnostics_txt or "").encode("utf-8")

    if receiving_date_from is not None and receiving_date_to is not None:
        period_text = format_report_period(receiving_date_from, receiving_date_to)
    else:
        period_text = ""

    new_wb = openpyxl.Workbook()
    new_ws = require_worksheet(new_wb)
    if template_ws.title:
        new_ws.title = template_ws.title

    _replace_sample_count_header_placeholders(template_ws, period_text, total_samples)
    current_row = _copy_template_header_rows(template_ws, new_ws, period_text, total_samples)
    template_table_top_row = data_rows[0][0]

    for branch_block in by_branch:
        branch_name = branch_block.branch_name or ""
        row_values = {r["label"]: r["value"] for r in branch_block.rows}

        block_start_row = current_row
        for template_row_idx, cell_b_value in data_rows:
            key = normalize_row_title_for_match(cell_b_value)
            if key and not is_sample_count_row_visible_for_branch(key, branch_name, ROW_TITLE_TO_BRANCH):
                continue
            value_c = match_row_title_to_value(cell_b_value, row_values)
            column_c_lines = _split_column_c_lines(value_c if value_c else SAMPLE_COUNT_EMPTY_CELL_VALUE)
            category_start_row = current_row
            category_row_height = _row_height_for_category(key or "", len(column_c_lines))

            for line_idx, line_text in enumerate(column_c_lines):
                out_row = current_row
                for col in (1, 2, 3):
                    src = template_ws.cell(row=template_row_idx, column=col)
                    tgt = new_ws.cell(row=out_row, column=col)
                    copy_cell_style(src, tgt)
                copy_row_formatting(
                    template_ws,
                    new_ws,
                    template_row_idx,
                    out_row,
                    merged_cells_map=None,
                )
                for col in (1, 2, 3):
                    report_cell = as_writable_cell(new_ws.cell(row=out_row, column=col))
                    if report_cell is not None:
                        apply_report_font(report_cell)
                _set_fixed_row_height(new_ws, out_row, category_row_height)

                set_cell_value(
                    new_ws,
                    out_row,
                    1,
                    branch_name if out_row == block_start_row else None,
                )
                if line_idx == 0:
                    set_cell_value(new_ws, out_row, 2, cell_b_value)

                cell_c = as_writable_cell(new_ws.cell(row=out_row, column=3))
                if cell_c is None:
                    current_row += 1
                    continue
                cell_c.value = line_text
                old_align = cell_c.alignment
                cell_c.alignment = Alignment(
                    wrap_text=False,
                    horizontal=getattr(old_align, "horizontal", "general"),
                    vertical=getattr(old_align, "vertical", "top"),
                    text_rotation=getattr(old_align, "text_rotation", 0),
                    shrink_to_fit=getattr(old_align, "shrink_to_fit", False),
                    indent=getattr(old_align, "indent", 0),
                )
                current_row += 1

            category_end_row = current_row - 1
            _apply_multiline_category_borders(new_ws, category_start_row, category_end_row)
            if category_end_row > category_start_row:
                new_ws.merge_cells(
                    start_row=category_start_row,
                    start_column=2,
                    end_row=category_end_row,
                    end_column=2,
                )
                cell_b = as_writable_cell(new_ws.cell(row=category_start_row, column=2))
                if cell_b is None:
                    continue
                old_b_align = cell_b.alignment
                cell_b.alignment = Alignment(
                    wrap_text=getattr(old_b_align, "wrap_text", False),
                    horizontal=getattr(old_b_align, "horizontal", "general"),
                    vertical="center",
                    text_rotation=getattr(old_b_align, "text_rotation", 0),
                    shrink_to_fit=getattr(old_b_align, "shrink_to_fit", False),
                    indent=getattr(old_b_align, "indent", 0),
                )

        block_end_row = current_row - 1
        if block_end_row >= block_start_row:
            if block_end_row > block_start_row:
                new_ws.merge_cells(
                    start_row=block_start_row,
                    start_column=1,
                    end_row=block_end_row,
                    end_column=1,
                )
            _apply_branch_block_top_border(
                template_ws,
                new_ws,
                block_start_row,
                template_table_top_row,
            )

    alignment_center = Alignment(horizontal="center", vertical="center")
    for row_idx in range(1, new_ws.max_row + 1):
        for col in (1, 2, 3):
            report_cell = as_writable_cell(new_ws.cell(row=row_idx, column=col))
            if report_cell is not None:
                apply_report_font(report_cell)
        align_cell = as_writable_cell(new_ws.cell(row=row_idx, column=1))
        if align_cell is not None:
            align_cell.alignment = alignment_center

    copy_column_dimensions(template_ws, new_ws)

    out = BytesIO()
    new_wb.save(out)
    return out.getvalue(), diagnostics_txt
