from __future__ import annotations
import base64
from io import BytesIO
import openpyxl
from openpyxl.worksheet.worksheet import Worksheet
import pendulum
from sqlalchemy.ext.asyncio import AsyncSession
from services.ilninm_reports.constants import (
    NKS_MAX_COLUMN,
    NKS_PLACEHOLDER_PERIOD,
    NKS_REPORT_FONT_SIZE,
    NKS_TEMPLATE_DATA_ROW,
    NKS_TEMPLATE_HEADER_LAST_ROW,
    REPORT_EMPTY_CELL_VALUE,
)
from services.ilninm_reports.nks import NksReportRow, format_nks_report_period, get_nks_report_rows
from utils.excel_typing import require_worksheet
from utils.ilninm_reports.excel_helpers import apply_report_font, merged_cell_anchor, writable_cell
from utils.ilninm_reports.excel_template_layout import (
    insert_rows_for_data_count,
    unmerge_cells_in_row_range,
)
from utils.protocol.template_markers import (
    copy_cell_style,
    copy_column_dimensions,
    copy_row_formatting,
)


def _template_data_row_height(ws: Worksheet, template_row: int) -> float | None:
    """Вернуть высоту образцовой строки данных в шаблоне."""
    row_dim = ws.row_dimensions.get(template_row)
    if row_dim and row_dim.height is not None:
        return row_dim.height
    return None


def _lock_data_row_height(
    ws: Worksheet,
    row: int,
    height: float | None,
) -> None:
    """Зафиксировать высоту строки данных по образцу шаблона."""
    if height is None:
        return
    ws.row_dimensions[row].height = height


def _replace_period_placeholder(
    ws: Worksheet,
    period_text: str,
) -> None:
    """Подставить период в метки шапки шаблона."""
    processed: set[tuple[int, int]] = set()
    max_col = max(ws.max_column, NKS_MAX_COLUMN)
    for row in range(1, NKS_TEMPLATE_HEADER_LAST_ROW + 1):
        for col in range(1, max_col + 1):
            anchor_row, anchor_col = merged_cell_anchor(ws, row, col)
            anchor = (anchor_row, anchor_col)
            if anchor in processed:
                continue
            processed.add(anchor)
            cell = ws.cell(row=anchor_row, column=anchor_col)
            if cell.value is None:
                continue
            text = str(cell.value)
            if NKS_PLACEHOLDER_PERIOD in text:
                cell.value = text.replace(NKS_PLACEHOLDER_PERIOD, period_text)


def _write_data_row(
    ws: Worksheet,
    template_row: int,
    output_row: int,
    row_index: int,
    well: str,
    sampling_date: str,
    values_by_column: dict[int, str],
    template_row_height: float | None,
) -> None:
    """Записать строку данных отчёта НКС на лист."""
    if output_row != template_row:
        copy_row_formatting(ws, ws, template_row, output_row, merged_cells_map=None)
        for col in range(1, NKS_MAX_COLUMN + 1):
            copy_cell_style(
                ws.cell(row=template_row, column=col),
                ws.cell(row=output_row, column=col),
            )

    writable_cell(ws, output_row, 1).value = row_index
    writable_cell(ws, output_row, 2).value = well
    writable_cell(ws, output_row, 3).value = sampling_date
    for col in range(4, NKS_MAX_COLUMN + 1):
        writable_cell(ws, output_row, col).value = values_by_column.get(col, REPORT_EMPTY_CELL_VALUE)

    for col in range(1, NKS_MAX_COLUMN + 1):
        apply_report_font(writable_cell(ws, output_row, col), size=NKS_REPORT_FONT_SIZE)

    _lock_data_row_height(ws, output_row, template_row_height)


async def build_nks_excel(
    db: AsyncSession,
    template_file_base64: str,
    laboratory_id: int,
    sampling_date_from: pendulum.DateTime,
    sampling_date_to: pendulum.DateTime,
    report_month: int,
    report_year: int,
    department_id: int | None = None,
) -> bytes:
    """Загрузить данные отчёта НКС и собрать Excel-файл."""
    rows = await get_nks_report_rows(
        db,
        laboratory_id,
        department_id,
        sampling_date_from,
        sampling_date_to,
    )
    return render_nks_excel(template_file_base64, report_month, report_year, rows)


def render_nks_excel(
    template_file_base64: str,
    report_month: int,
    report_year: int,
    rows: list[NksReportRow],
) -> bytes:
    """Собрать байты Excel-отчёта НКС по шаблону и готовым строкам."""
    template_bytes = base64.b64decode(template_file_base64)
    wb = openpyxl.load_workbook(BytesIO(template_bytes))
    ws = wb.worksheets[0] if wb.worksheets else require_worksheet(wb)

    period_text = format_nks_report_period(report_month, report_year)
    for sheet in wb.worksheets:
        _replace_period_placeholder(sheet, period_text)

    template_row = NKS_TEMPLATE_DATA_ROW
    template_row_height = _template_data_row_height(ws, template_row)
    data_row_count = len(rows)
    insert_rows_for_data_count(ws, template_row, data_row_count)
    if data_row_count > 0:
        unmerge_cells_in_row_range(
            ws,
            template_row,
            template_row + data_row_count - 1,
            NKS_MAX_COLUMN,
        )
    output_row = template_row

    for data_row in rows:
        _write_data_row(
            ws,
            template_row,
            output_row,
            data_row.row_index,
            data_row.well,
            data_row.sampling_date,
            data_row.values_by_column,
            template_row_height,
        )
        output_row += 1

    copy_column_dimensions(ws, ws)

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
