import base64
from io import BytesIO
from typing import Any, Optional
import openpyxl
from openpyxl.cell.cell import Cell
from openpyxl.styles import Font
from sqlalchemy.ext.asyncio import AsyncSession
from services.ilninm_reports.constants import (
    NKS_MAX_COLUMN,
    NKS_PLACEHOLDER_PERIOD,
    NKS_REPORT_FONT_SIZE,
    NKS_TEMPLATE_DATA_ROW,
    NKS_TEMPLATE_HEADER_LAST_ROW,
    REPORT_EMPTY_CELL_VALUE,
)
from services.ilninm_reports.excel_template_layout import (
    insert_rows_for_data_count,
    unmerge_cells_in_row_range,
)
from services.ilninm_reports.nks import format_nks_report_period, get_nks_report_rows
from services.ilninm_reports.sample_count_generator import REPORT_FONT_NAME
from utils.protocol_generator_utils import (
    copy_cell_style,
    copy_column_dimensions,
    copy_row_formatting,
)


def _apply_report_font(cell: Cell) -> None:
    old = cell.font
    if old:
        cell.font = Font(
            name=REPORT_FONT_NAME,
            size=NKS_REPORT_FONT_SIZE,
            bold=old.bold,
            italic=old.italic,
            underline=old.underline,
            strike=old.strike,
            color="FF000000",
        )
        return
    cell.font = Font(name=REPORT_FONT_NAME, size=NKS_REPORT_FONT_SIZE, color="FF000000")


def _merged_cell_anchor(
    ws: openpyxl.worksheet.worksheet.Worksheet, row: int, col: int
) -> tuple[int, int]:
    for merged in ws.merged_cells.ranges:
        if (
            merged.min_row <= row <= merged.max_row
            and merged.min_col <= col <= merged.max_col
        ):
            return merged.min_row, merged.min_col
    return row, col


def _writable_cell(
    ws: openpyxl.worksheet.worksheet.Worksheet, row: int, col: int
) -> Cell:
    anchor_row, anchor_col = _merged_cell_anchor(ws, row, col)
    return ws.cell(row=anchor_row, column=anchor_col)


def _template_data_row_height(
    ws: openpyxl.worksheet.worksheet.Worksheet, template_row: int
) -> Optional[float]:
    row_dim = ws.row_dimensions.get(template_row)
    if row_dim and row_dim.height is not None:
        return row_dim.height
    return None


def _lock_data_row_height(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    row: int,
    height: Optional[float],
) -> None:
    if height is None:
        return
    ws.row_dimensions[row].height = height


def _replace_period_placeholder(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    period_text: str,
) -> None:
    processed: set[tuple[int, int]] = set()
    max_col = max(ws.max_column, NKS_MAX_COLUMN)
    for row in range(1, NKS_TEMPLATE_HEADER_LAST_ROW + 1):
        for col in range(1, max_col + 1):
            anchor_row, anchor_col = _merged_cell_anchor(ws, row, col)
            anchor = (anchor_row, anchor_col)
            if anchor in processed:
                continue
            processed.add(anchor)
            cell = ws.cell(row=anchor_row, column=anchor_col)
            if cell.value is None:
                continue
            text = str(cell.value)
            if NKS_PLACEHOLDER_PERIOD in text:
                cell.value = text.replace(NKS_PLACEHOLDER_PERIOD, period_text)


def _write_data_row(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    template_row: int,
    output_row: int,
    row_index: int,
    well: str,
    sampling_date: str,
    values_by_column: dict[int, str],
    template_row_height: Optional[float],
) -> None:
    if output_row != template_row:
        copy_row_formatting(ws, ws, template_row, output_row, merged_cells_map=None)
        for col in range(1, NKS_MAX_COLUMN + 1):
            copy_cell_style(
                ws.cell(row=template_row, column=col),
                ws.cell(row=output_row, column=col),
            )

    _writable_cell(ws, output_row, 1).value = row_index
    _writable_cell(ws, output_row, 2).value = well
    _writable_cell(ws, output_row, 3).value = sampling_date
    for col in range(4, NKS_MAX_COLUMN + 1):
        _writable_cell(ws, output_row, col).value = values_by_column.get(
            col, REPORT_EMPTY_CELL_VALUE
        )

    for col in range(1, NKS_MAX_COLUMN + 1):
        _apply_report_font(_writable_cell(ws, output_row, col))

    _lock_data_row_height(ws, output_row, template_row_height)


async def build_nks_excel(
    db: AsyncSession,
    template_file_base64: str,
    laboratory_id: int,
    sampling_date_from: Any,
    sampling_date_to: Any,
    report_month: int,
    report_year: int,
    department_id: Optional[int] = None,
) -> bytes:
    """Строит Excel по шаблону: шапка с периодом, строки данных с 19-й строки."""
    template_bytes = base64.b64decode(template_file_base64)
    wb = openpyxl.load_workbook(BytesIO(template_bytes))
    ws = wb.worksheets[0] if wb.worksheets else wb.active

    period_text = format_nks_report_period(report_month, report_year)
    for sheet in wb.worksheets:
        _replace_period_placeholder(sheet, period_text)

    rows = await get_nks_report_rows(
        db,
        laboratory_id,
        department_id,
        sampling_date_from,
        sampling_date_to,
    )

    template_row = NKS_TEMPLATE_DATA_ROW
    template_row_height = _template_data_row_height(ws, template_row)
    data_row_count = len(rows)
    insert_rows_for_data_count(ws, template_row, data_row_count)
    if data_row_count > 0:
        unmerge_cells_in_row_range(
            ws,
            template_row,
            template_row + data_row_count - 1,
            NKS_MAX_COLUMN,
        )
    output_row = template_row

    for data_row in rows:
        _write_data_row(
            ws,
            template_row,
            output_row,
            data_row.row_index,
            data_row.well,
            data_row.sampling_date,
            data_row.values_by_column,
            template_row_height,
        )
        output_row += 1

    copy_column_dimensions(ws, ws)

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
