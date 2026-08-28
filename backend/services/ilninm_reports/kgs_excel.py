from __future__ import annotations
import base64
from io import BytesIO
import openpyxl
from openpyxl.worksheet.worksheet import Worksheet
import pendulum
from sqlalchemy.ext.asyncio import AsyncSession
from services.ilninm_reports.constants import (
    KGS_AVERAGE_ROW_LABEL,
    KGS_PLACEHOLDER_PERIOD,
    KGS_TEMPLATE_DATA_ROW,
    KGS_TEMPLATE_HEADER_LAST_ROW,
    REPORT_EMPTY_CELL_VALUE,
)
from services.ilninm_reports.kgs import (
    KGS_METHOD_COLUMNS,
    KgsReportLocationGroup,
    format_report_period,
    get_kgs_report_groups,
)
from utils.excel_typing import require_worksheet
from utils.ilninm_reports.excel_helpers import apply_report_font, merged_cell_anchor, writable_cell
from utils.ilninm_reports.excel_template_layout import (
    insert_rows_for_data_count,
    unmerge_cells_in_row_range,
)
from utils.protocol.template_markers import (
    copy_cell_style,
    copy_column_dimensions,
    copy_row_formatting,
)


def _replace_period_placeholder(
    ws: Worksheet,
    period_text: str,
) -> None:
    """Подставить период в метки шапки шаблона."""
    processed: set[tuple[int, int]] = set()
    max_col = max(ws.max_column, 8)
    for row in range(1, KGS_TEMPLATE_HEADER_LAST_ROW + 1):
        for col in range(1, max_col + 1):
            anchor_row, anchor_col = merged_cell_anchor(ws, row, col)
            anchor = (anchor_row, anchor_col)
            if anchor in processed:
                continue
            processed.add(anchor)
            cell = ws.cell(row=anchor_row, column=anchor_col)
            if cell.value is None:
                continue
            text = str(cell.value)
            if KGS_PLACEHOLDER_PERIOD in text:
                cell.value = text.replace(KGS_PLACEHOLDER_PERIOD, period_text)


def _write_data_row(
    ws: Worksheet,
    template_row: int,
    output_row: int,
    max_col: int,
    row_index: int | None,
    location_display: str,
    sampling_date: str,
    values_by_column: dict[int, str],
) -> None:
    """Записать строку данных отчёта КГС на лист."""
    if output_row != template_row:
        copy_row_formatting(ws, ws, template_row, output_row, merged_cells_map=None)
        for col in range(1, max_col + 1):
            copy_cell_style(
                ws.cell(row=template_row, column=col),
                ws.cell(row=output_row, column=col),
            )

    writable_cell(ws, output_row, 1).value = row_index
    writable_cell(ws, output_row, 2).value = location_display
    writable_cell(ws, output_row, 3).value = sampling_date
    for spec in KGS_METHOD_COLUMNS:
        writable_cell(ws, output_row, spec.column).value = values_by_column.get(spec.column, REPORT_EMPTY_CELL_VALUE)

    for col in range(1, max_col + 1):
        apply_report_font(writable_cell(ws, output_row, col))


def _write_average_row(
    ws: Worksheet,
    template_row: int,
    output_row: int,
    max_col: int,
    values_by_column: dict[int, str],
) -> None:
    """Записать строку средних значений по месту отбора."""
    if output_row != template_row:
        copy_row_formatting(ws, ws, template_row, output_row, merged_cells_map=None)
        for col in range(1, max_col + 1):
            copy_cell_style(
                ws.cell(row=template_row, column=col),
                ws.cell(row=output_row, column=col),
            )

    writable_cell(ws, output_row, 1).value = None
    writable_cell(ws, output_row, 2).value = KGS_AVERAGE_ROW_LABEL
    writable_cell(ws, output_row, 3).value = None
    for spec in KGS_METHOD_COLUMNS:
        writable_cell(ws, output_row, spec.column).value = values_by_column.get(spec.column, REPORT_EMPTY_CELL_VALUE)

    for col in range(1, max_col + 1):
        apply_report_font(writable_cell(ws, output_row, col))


async def build_kgs_excel(
    db: AsyncSession,
    template_file_base64: str,
    laboratory_id: int,
    sampling_date_from: pendulum.DateTime,
    sampling_date_to: pendulum.DateTime,
    department_id: int | None = None,
) -> bytes:
    """Загрузить данные отчёта КГС и собрать Excel-файл."""
    groups = await get_kgs_report_groups(
        db,
        laboratory_id,
        department_id,
        sampling_date_from,
        sampling_date_to,
    )
    return render_kgs_excel(
        template_file_base64,
        sampling_date_from,
        sampling_date_to,
        groups,
    )


def render_kgs_excel(
    template_file_base64: str,
    sampling_date_from: pendulum.DateTime,
    sampling_date_to: pendulum.DateTime,
    groups: list[KgsReportLocationGroup],
) -> bytes:
    """Собрать байты Excel-отчёта КГС по шаблону и готовым группам данных."""
    template_bytes = base64.b64decode(template_file_base64)
    wb = openpyxl.load_workbook(BytesIO(template_bytes))
    ws = wb.worksheets[0] if wb.worksheets else require_worksheet(wb)

    period_text = format_report_period(sampling_date_from, sampling_date_to)
    for sheet in wb.worksheets:
        _replace_period_placeholder(sheet, period_text)

    template_row = KGS_TEMPLATE_DATA_ROW
    max_col = max((spec.column for spec in KGS_METHOD_COLUMNS), default=8)
    data_row_count = sum(len(group.data_rows) + 1 for group in groups)
    insert_rows_for_data_count(ws, template_row, data_row_count)
    if data_row_count > 0:
        unmerge_cells_in_row_range(
            ws,
            template_row,
            template_row + data_row_count - 1,
            max_col,
        )
    output_row = template_row
    row_index = 1

    for group in groups:
        for data_row in group.data_rows:
            _write_data_row(
                ws,
                template_row,
                output_row,
                max_col,
                row_index,
                data_row.location_display,
                data_row.sampling_date,
                data_row.values_by_column,
            )
            row_index += 1
            output_row += 1

        _write_average_row(
            ws,
            template_row,
            output_row,
            max_col,
            group.average_row.values_by_column,
        )
        output_row += 1

    copy_column_dimensions(ws, ws)

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
