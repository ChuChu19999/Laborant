import base64
from copy import copy
from io import BytesIO
from typing import Any
import openpyxl
import pendulum
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, Side
from sqlalchemy.ext.asyncio import AsyncSession
from services.ilninm_reports.constants import (
    ROW_TITLE_KALIBROVOCHNAYA_NEFT_UGPU,
    ROW_TITLE_TOVARNAYA_NEFT_NGDU,
    SAMPLE_COUNT_PLACEHOLDER_KOL_VO,
    SAMPLE_COUNT_PLACEHOLDER_PERIOD,
    SAMPLE_COUNT_TEMPLATE_HEADER_ROW_COUNT,
)
from services.ilninm_reports.physicochemical import format_report_period
from services.ilninm_reports.sample_count import (
    ROW_TITLE_TO_BRANCH,
    _normalize_cell_a_for_match,
    get_sample_count_report_data,
    is_sample_count_row_visible_for_branch,
    match_row_title_to_value,
)
from utils.protocol_generator_utils import (
    copy_cell_style,
    copy_column_dimensions,
    copy_row_formatting,
)

# Шрифт всего отчёта.
REPORT_FONT_NAME = "Times New Roman"
REPORT_FONT_SIZE = 12

# Высота строки с текстом в столбце C (пункты Excel).
ROW_HEIGHT_COLUMN_C_LINES = 15
ROW_HEIGHT_COLUMN_C_SINGLE_LINE_TALL = 30

_ROW_TITLES_TALL_WHEN_SINGLE_LINE = frozenset(
    {
        ROW_TITLE_TOVARNAYA_NEFT_NGDU.lower(),
        ROW_TITLE_KALIBROVOCHNAYA_NEFT_UGPU.lower(),
    }
)


def _split_column_c_lines(value: str) -> list[str]:
    """Разбивает значение столбца C на строки листа; хвостовые пустые строки отбрасываются."""
    if not value or value == "—":
        return ["—"]
    parts = value.split("\n")
    while parts and not parts[-1].strip():
        parts.pop()
    return parts if parts else ["—"]


def _row_height_for_category(category_key: str, lines_in_category: int) -> float:
    """30 pt для однострочных «Товарная нефть НГДУ» и «Калибровочная нефть УГПУ», иначе 15 pt."""
    if (
        lines_in_category == 1
        and category_key.lower() in _ROW_TITLES_TALL_WHEN_SINGLE_LINE
    ):
        return ROW_HEIGHT_COLUMN_C_SINGLE_LINE_TALL
    return ROW_HEIGHT_COLUMN_C_LINES


def _set_fixed_row_height(
    ws: openpyxl.worksheet.worksheet.Worksheet, row: int, height: float
) -> None:
    """Фиксированная высота строки для вывода текста в C."""
    ws.row_dimensions[row].height = height


def _apply_report_font(cell: Cell) -> None:
    """Times New Roman 12, черный цвет; начертание из шаблона сохраняется."""
    old = cell.font
    if old:
        cell.font = Font(
            name=REPORT_FONT_NAME,
            size=REPORT_FONT_SIZE,
            bold=old.bold,
            italic=old.italic,
            underline=old.underline,
            strike=old.strike,
            color="FF000000",
        )
        return
    cell.font = Font(name=REPORT_FONT_NAME, size=REPORT_FONT_SIZE, color="FF000000")


def _cell_border_without_horizontal_edges(
    cell: Cell,
    *,
    remove_top: bool,
    remove_bottom: bool,
) -> None:
    """Снимает верхнюю и/или нижнюю границу ячейки, боковые не трогает."""
    if not remove_top and not remove_bottom:
        return
    old = cell.border or Border()
    cell.border = Border(
        left=old.left,
        right=old.right,
        top=Side(style=None) if remove_top else old.top,
        bottom=Side(style=None) if remove_bottom else old.bottom,
    )


def _apply_multiline_category_borders(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    start_row: int,
    end_row: int,
) -> None:
    """
    Между строками одной категории в C убирает горизонтальные границы:
    у первой строки — нижнюю, у последней — верхнюю, у средних — обе.
    """
    line_count = end_row - start_row + 1
    if line_count <= 1:
        return
    for idx, row in enumerate(range(start_row, end_row + 1)):
        _cell_border_without_horizontal_edges(
            ws.cell(row=row, column=3),
            remove_top=idx > 0,
            remove_bottom=idx < line_count - 1,
        )


def _merged_cell_anchor(
    ws: openpyxl.worksheet.worksheet.Worksheet, row: int, col: int
) -> tuple[int, int]:
    for merged in ws.merged_cells.ranges:
        if (
            merged.min_row <= row <= merged.max_row
            and merged.min_col <= col <= merged.max_col
        ):
            return merged.min_row, merged.min_col
    return row, col


def _copy_worksheet_merged_ranges(
    src_ws: openpyxl.worksheet.worksheet.Worksheet,
    tgt_ws: openpyxl.worksheet.worksheet.Worksheet,
    *,
    min_row: int,
    max_row: int,
    row_offset: int = 0,
) -> None:
    """Копирует объединения ячеек из шаблона, полностью попадающие в диапазон строк."""
    for merged_range in list(src_ws.merged_cells.ranges):
        if merged_range.min_row < min_row or merged_range.max_row > max_row:
            continue
        tgt_ws.merge_cells(
            start_row=merged_range.min_row + row_offset,
            start_column=merged_range.min_col,
            end_row=merged_range.max_row + row_offset,
            end_column=merged_range.max_col,
        )


def _apply_cell_top_border_from(source: Cell, target: Cell) -> None:
    """Верхняя граница из образца (для верхней строки блока филиала)."""
    src_border = source.border
    if not src_border or not src_border.top or src_border.top.style is None:
        return
    old = target.border or Border()
    target.border = Border(
        left=old.left,
        right=old.right,
        top=copy(src_border.top),
        bottom=old.bottom,
    )


def _apply_branch_block_top_border(
    template_ws: openpyxl.worksheet.worksheet.Worksheet,
    new_ws: openpyxl.worksheet.worksheet.Worksheet,
    block_start_row: int,
    template_top_row: int,
) -> None:
    """Верхняя граница рамки блока филиала — как у первой строки категорий в шаблоне."""
    for col in (1, 2, 3):
        _apply_cell_top_border_from(
            template_ws.cell(row=template_top_row, column=col),
            new_ws.cell(row=block_start_row, column=col),
        )


def _replace_text_placeholders(text: str, period_text: str, total_samples: int) -> str:
    updated = text
    if SAMPLE_COUNT_PLACEHOLDER_PERIOD in updated:
        updated = updated.replace(SAMPLE_COUNT_PLACEHOLDER_PERIOD, period_text)
    if SAMPLE_COUNT_PLACEHOLDER_KOL_VO in updated:
        updated = updated.replace(SAMPLE_COUNT_PLACEHOLDER_KOL_VO, str(total_samples))
    return updated


def _replace_sample_count_header_placeholders(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    period_text: str,
    total_samples: int,
) -> None:
    """Подставляет период и общее число проб в метки шапки шаблона."""
    processed: set[tuple[int, int]] = set()
    max_col = max(ws.max_column, 3)
    for row in range(1, SAMPLE_COUNT_TEMPLATE_HEADER_ROW_COUNT + 1):
        for col in range(1, max_col + 1):
            anchor_row, anchor_col = _merged_cell_anchor(ws, row, col)
            anchor = (anchor_row, anchor_col)
            if anchor in processed:
                continue
            processed.add(anchor)
            cell = ws.cell(row=anchor_row, column=anchor_col)
            if cell.value is None:
                continue
            text = str(cell.value)
            replaced = _replace_text_placeholders(text, period_text, total_samples)
            if replaced != text:
                cell.value = replaced


def _copy_template_header_rows(
    template_ws: openpyxl.worksheet.worksheet.Worksheet,
    new_ws: openpyxl.worksheet.worksheet.Worksheet,
    period_text: str,
    total_samples: int,
) -> int:
    """
    Копирует первые строки шаблона (шапка) в новый лист, подставляет метки.
    Возвращает номер следующей свободной строки на листе.
    """
    max_col = max(template_ws.max_column, 3)
    for row in range(1, SAMPLE_COUNT_TEMPLATE_HEADER_ROW_COUNT + 1):
        copy_row_formatting(template_ws, new_ws, row, row, merged_cells_map=None)
        for col in range(1, max_col + 1):
            anchor_row, anchor_col = _merged_cell_anchor(template_ws, row, col)
            src = template_ws.cell(row=anchor_row, column=anchor_col)
            tgt = new_ws.cell(row=row, column=col)
            copy_cell_style(src, tgt)
            if row == anchor_row and col == anchor_col and src.value is not None:
                tgt.value = _replace_text_placeholders(
                    str(src.value), period_text, total_samples
                )
        if template_ws.row_dimensions[row].height is not None:
            new_ws.row_dimensions[row].height = template_ws.row_dimensions[row].height
    _copy_worksheet_merged_ranges(
        template_ws,
        new_ws,
        min_row=1,
        max_row=SAMPLE_COUNT_TEMPLATE_HEADER_ROW_COUNT,
    )
    return SAMPLE_COUNT_TEMPLATE_HEADER_ROW_COUNT + 1


def _get_cell_b_value(ws: openpyxl.worksheet.worksheet.Worksheet, row: int) -> Any:
    """Значение столбца B с учётом объединённых ячеек (берём верхнюю ячейку слияния)."""
    cell = ws.cell(row=row, column=2)
    if cell.value is not None:
        return cell.value
    for merged in ws.merged_cells.ranges:
        if (
            merged.min_col <= 2 <= merged.max_col
            and merged.min_row <= row <= merged.max_row
        ):
            return ws.cell(row=merged.min_row, column=2).value
    return None


def _find_template_data_rows(
    ws: openpyxl.worksheet.worksheet.Worksheet,
) -> list[tuple[int, Any]]:
    """
    Находит в шаблоне строки с категориями в столбце B.
    Учитывает объединённые ячейки: для каждой строки блока возвращается одна запись,
    чтобы последняя строка блока не терялась. Возвращает список (номер_строки, значение_ячейки_B).
    """
    known_titles = {
        "Товарная нефть НГДУ",
        "Эксплуатационная нефть НГДУ",
        "Калибровочная нефть УГПУ",
        "Внеплановые",
        "Паспортизация",
        "ГКП-21 ГКП-22",
        "ОИС Ачимовка",
        "ОИС Валанжин",
        "ОИС Ен-Яха",
        "ОИС",
        "Товарная продукция ОИС",
        "Прочие",
        "Нефтеконденсатная смесь",
        "Дизтопливо",
        "Ингибитор коррозии",
    }
    result = []
    for row_idx in range(1, ws.max_row + 1):
        val = _get_cell_b_value(ws, row_idx)
        key = _normalize_cell_a_for_match(val)
        if not key:
            continue
        key_lower = key.lower()
        for title in known_titles:
            if title.lower() == key_lower:
                result.append((row_idx, val))
                break
    if not result:
        return result
    min_row = result[0][0]
    max_row = result[-1][0]
    expanded = []
    for merged in ws.merged_cells.ranges:
        if merged.min_col > 2 or merged.max_col < 2:
            continue
        if merged.max_row <= min_row or merged.min_row >= max_row:
            continue
        top_val = ws.cell(row=merged.min_row, column=2).value
        if top_val is None:
            continue
        for r in range(merged.min_row, merged.max_row + 1):
            if any(t[0] == r for t in result):
                continue
            expanded.append((r, top_val))
    result = sorted(result + expanded, key=lambda x: x[0])
    return result


async def build_sample_count_excel(
    db: AsyncSession,
    template_file_base64: str,
    laboratory_id: int,
    receiving_date_from: Any | None,
    receiving_date_to: Any | None,
    department_id: int | None = None,
) -> tuple[bytes, bytes]:
    """
    Строит Excel-файл отчёта «Количество проб»: для каждого branch копируется
    блок шаблона (категории в столбце B), в столбец C подставляются данные.
    В первую строку каждого блока в столбец A записывается наименование branch.
    Пробы фильтруются по laboratory_id и department_id.
    """
    template_bytes = BytesIO(base64.b64decode(template_file_base64))
    template_wb = openpyxl.load_workbook(template_bytes)
    template_ws = template_wb.active
    data_rows = _find_template_data_rows(template_ws)
    if not data_rows:
        empty_txt = (
            "Диагностика не сформирована: в шаблоне не найдены строки категорий.\n"
        )
        return template_bytes.getvalue(), empty_txt.encode("utf-8")

    report_data = await get_sample_count_report_data(
        db,
        laboratory_id=laboratory_id,
        receiving_date_from=receiving_date_from,
        receiving_date_to=receiving_date_to,
        department_id=department_id,
    )
    by_branch = report_data.get("by_branch") or []
    total_samples = int(report_data.get("total_samples") or 0)
    diagnostics_txt = (report_data.get("diagnostics_txt") or "").encode("utf-8")

    if receiving_date_from is not None and receiving_date_to is not None:
        period_text = format_report_period(
            pendulum.instance(receiving_date_from),
            pendulum.instance(receiving_date_to),
        )
    else:
        period_text = ""

    new_wb = openpyxl.Workbook()
    new_ws = new_wb.active
    if template_ws.title:
        new_ws.title = template_ws.title

    _replace_sample_count_header_placeholders(template_ws, period_text, total_samples)
    current_row = _copy_template_header_rows(
        template_ws, new_ws, period_text, total_samples
    )
    template_table_top_row = data_rows[0][0]

    for branch_block in by_branch:
        branch_name = branch_block.get("branch_name") or ""
        rows = branch_block.get("rows") or []
        row_values = {r["label"]: r["value"] for r in rows}

        block_start_row = current_row
        for template_row_idx, cell_b_value in data_rows:
            key = _normalize_cell_a_for_match(cell_b_value)
            if key and not is_sample_count_row_visible_for_branch(
                key, branch_name, ROW_TITLE_TO_BRANCH
            ):
                continue
            value_c = match_row_title_to_value(cell_b_value, row_values)
            column_c_lines = _split_column_c_lines(value_c if value_c else "—")
            category_start_row = current_row
            category_row_height = _row_height_for_category(
                key or "", len(column_c_lines)
            )

            for line_idx, line_text in enumerate(column_c_lines):
                out_row = current_row
                for col in (1, 2, 3):
                    src = template_ws.cell(row=template_row_idx, column=col)
                    tgt = new_ws.cell(row=out_row, column=col)
                    copy_cell_style(src, tgt)
                copy_row_formatting(
                    template_ws,
                    new_ws,
                    template_row_idx,
                    out_row,
                    merged_cells_map=None,
                )
                for col in (1, 2, 3):
                    _apply_report_font(new_ws.cell(row=out_row, column=col))
                _set_fixed_row_height(new_ws, out_row, category_row_height)

                new_ws.cell(row=out_row, column=1).value = (
                    branch_name if out_row == block_start_row else None
                )
                if line_idx == 0:
                    new_ws.cell(row=out_row, column=2).value = cell_b_value

                cell_c = new_ws.cell(row=out_row, column=3)
                cell_c.value = line_text
                old_align = cell_c.alignment
                cell_c.alignment = Alignment(
                    wrap_text=False,
                    horizontal=getattr(old_align, "horizontal", "general"),
                    vertical=getattr(old_align, "vertical", "top"),
                    text_rotation=getattr(old_align, "text_rotation", 0),
                    shrink_to_fit=getattr(old_align, "shrink_to_fit", False),
                    indent=getattr(old_align, "indent", 0),
                )
                current_row += 1

            category_end_row = current_row - 1
            _apply_multiline_category_borders(
                new_ws, category_start_row, category_end_row
            )
            if category_end_row > category_start_row:
                new_ws.merge_cells(
                    start_row=category_start_row,
                    start_column=2,
                    end_row=category_end_row,
                    end_column=2,
                )
                cell_b = new_ws.cell(row=category_start_row, column=2)
                old_b_align = cell_b.alignment
                cell_b.alignment = Alignment(
                    wrap_text=getattr(old_b_align, "wrap_text", False),
                    horizontal=getattr(old_b_align, "horizontal", "general"),
                    vertical="center",
                    text_rotation=getattr(old_b_align, "text_rotation", 0),
                    shrink_to_fit=getattr(old_b_align, "shrink_to_fit", False),
                    indent=getattr(old_b_align, "indent", 0),
                )

        block_end_row = current_row - 1
        if block_end_row >= block_start_row:
            if block_end_row > block_start_row:
                new_ws.merge_cells(
                    start_row=block_start_row,
                    start_column=1,
                    end_row=block_end_row,
                    end_column=1,
                )
            _apply_branch_block_top_border(
                template_ws,
                new_ws,
                block_start_row,
                template_table_top_row,
            )

    alignment_center = Alignment(horizontal="center", vertical="center")
    for row_idx in range(1, new_ws.max_row + 1):
        for col in (1, 2, 3):
            _apply_report_font(new_ws.cell(row=row_idx, column=col))
        new_ws.cell(row=row_idx, column=1).alignment = alignment_center

    copy_column_dimensions(template_ws, new_ws)

    out = BytesIO()
    new_wb.save(out)
    return out.getvalue(), diagnostics_txt
