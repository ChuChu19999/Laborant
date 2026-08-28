from __future__ import annotations
import base64
from io import BytesIO
import openpyxl
from openpyxl.cell.cell import Cell
from openpyxl.worksheet.worksheet import Worksheet
import pendulum
from sqlalchemy.ext.asyncio import AsyncSession
from services.ilninm_reports.constants import (
    PHYSICOCHEMICAL_PLACEHOLDER_PERIOD,
    PHYSICOCHEMICAL_PLACEHOLDER_SAMPLING_LOCATION,
    PHYSICOCHEMICAL_TEMPLATE_DATA_ROW,
    PHYSICOCHEMICAL_TEMPLATE_HEADER_LAST_ROW,
    REPORT_EMPTY_CELL_VALUE,
)
from services.ilninm_reports.physicochemical import (
    METHOD_COLUMNS,
    PhysicochemicalReportRow,
    format_report_period,
    get_physicochemical_report_rows,
    resolve_sampling_location_display,
)
from utils.excel_typing import as_writable_cell, require_worksheet, set_cell_value
from utils.ilninm_reports.excel_helpers import apply_report_font, merged_cell_anchor
from utils.protocol.template_markers import (
    copy_cell_style,
    copy_column_dimensions,
    copy_row_formatting,
)


def _apply_placeholders_to_cell(
    cell: Cell,
    period_text: str,
    sampling_location_text: str,
) -> None:
    """Подставить период и место отбора в текст ячейки шапки."""
    if cell.value is None:
        return
    text = str(cell.value)
    updated = text
    if PHYSICOCHEMICAL_PLACEHOLDER_PERIOD in updated:
        updated = updated.replace(PHYSICOCHEMICAL_PLACEHOLDER_PERIOD, period_text)
    if PHYSICOCHEMICAL_PLACEHOLDER_SAMPLING_LOCATION in updated:
        updated = updated.replace(
            PHYSICOCHEMICAL_PLACEHOLDER_SAMPLING_LOCATION,
            sampling_location_text,
        )
    if updated != text:
        cell.value = updated


def _replace_header_placeholders(
    ws: Worksheet,
    period_text: str,
    sampling_location_text: str,
) -> None:
    """Подставить период и место отбора в шапку шаблона."""
    processed: set[tuple[int, int]] = set()
    max_col = max(ws.max_column, 30)
    for row in range(1, PHYSICOCHEMICAL_TEMPLATE_HEADER_LAST_ROW + 1):
        for col in range(1, max_col + 1):
            anchor_row, anchor_col = merged_cell_anchor(ws, row, col)
            anchor = (anchor_row, anchor_col)
            if anchor in processed:
                continue
            processed.add(anchor)
            cell = as_writable_cell(ws.cell(row=anchor_row, column=anchor_col))
            if cell is None:
                continue
            _apply_placeholders_to_cell(
                cell,
                period_text,
                sampling_location_text,
            )


async def build_physicochemical_excel(
    db: AsyncSession,
    template_file_base64: str,
    laboratory_id: int,
    sampling_date_from: pendulum.DateTime,
    sampling_date_to: pendulum.DateTime,
    sampling_location: str,
    department_id: int | None = None,
) -> bytes:
    """Загрузить данные физико-химического отчёта и собрать Excel-файл."""
    rows = await get_physicochemical_report_rows(
        db,
        laboratory_id,
        department_id,
        sampling_date_from,
        sampling_date_to,
        sampling_location,
    )
    return render_physicochemical_excel(
        template_file_base64,
        sampling_date_from,
        sampling_date_to,
        sampling_location,
        rows,
    )


def render_physicochemical_excel(
    template_file_base64: str,
    sampling_date_from: pendulum.DateTime,
    sampling_date_to: pendulum.DateTime,
    sampling_location: str,
    rows: list[PhysicochemicalReportRow],
) -> bytes:
    """Собрать байты Excel по шаблону и готовым строкам отчёта."""
    template_bytes = base64.b64decode(template_file_base64)
    wb = openpyxl.load_workbook(BytesIO(template_bytes))
    ws = wb.worksheets[0] if wb.worksheets else require_worksheet(wb)

    period_text = format_report_period(sampling_date_from, sampling_date_to)
    location_text = resolve_sampling_location_display(sampling_location)
    for sheet in wb.worksheets:
        _replace_header_placeholders(sheet, period_text, location_text)

    template_row = PHYSICOCHEMICAL_TEMPLATE_DATA_ROW
    max_col = max((spec.column for spec in METHOD_COLUMNS), default=19)
    output_row = template_row

    for index, row_data in enumerate(rows, start=1):
        if output_row != template_row:
            copy_row_formatting(ws, ws, template_row, output_row, merged_cells_map=None)
            for col in range(1, max_col + 1):
                src = ws.cell(row=template_row, column=col)
                tgt = ws.cell(row=output_row, column=col)
                copy_cell_style(src, tgt)

        set_cell_value(ws, output_row, 1, index)
        set_cell_value(ws, output_row, 2, row_data.well)
        set_cell_value(ws, output_row, 3, row_data.sampling_date)
        for spec in METHOD_COLUMNS:
            set_cell_value(
                ws,
                output_row,
                spec.column,
                row_data.values_by_column.get(spec.column, REPORT_EMPTY_CELL_VALUE),
            )

        for col in range(1, max_col + 1):
            report_cell = as_writable_cell(ws.cell(row=output_row, column=col))
            if report_cell is not None:
                apply_report_font(report_cell)

        output_row += 1

    copy_column_dimensions(ws, ws)

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
